from __future__ import annotations

import csv
import io
from collections import defaultdict

from .units import normalize_assay, to_wt_percent, W_HEADERS
from .domain import finite_number, parse_elements, recipe, validate_sample, ELEMENTS
from .precursors import COLUMNS as PRECURSOR_COLUMNS, from_row as precursor_from_row, to_columns as precursor_to_columns, text as precursor_text

SAMPLE_COLUMNS = ["sample_id", "batch_id", "origin", "target_elements", "pulse_count", "temperature_C", "pulse_limit_C", "voltage_V", "current_A", "precursor_family", "support_batch", "protocol_id", "xrd_state", "xrd_source", "xrd_criteria", "xrd_reviewer", "stem_state", "stem_source", "stem_reviewer", "stem_notes", "notes", "icp_state", "icp_source", "icp_reviewer", "xrd_mode"]
SAMPLE_COLUMNS += ["recommendation_id", "planned_temperature_C", "sample_condition", "record_uid", "pan_mass_g", "solvent_volume_mL", "preparation_basis", "concentration_input_unit"]
SAMPLE_COLUMNS += ['electrical_input_kind']
ICP_COLUMNS = ["sample_id", "element", "w_wt_percent", "unit", "flag", "method_id", "source", "precursor_mM"]
ICP_COLUMNS += PRECURSOR_COLUMNS
METHOD_COLUMNS = ["method_id", "origin", "instrument", "rule", "element", "lod_wt_percent", "loq_wt_percent", "confirmed", "source", "version"]


def read_table(content, name, sheet=None):
    if name.lower().endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            if sheet not in book.sheetnames:
                raise ValueError(f"Excel must include the standard worksheet: {sheet}; found {book.sheetnames}. Instrument-specific report adapters are not connected yet.")
            rows = list(book[sheet].values)
        finally:
            book.close()
        if not rows:
            return []
        headers = [str(v or "").strip() for v in rows[0]]
        named = [h for h in headers if h]
        if len(named) != len(set(named)):
            raise ValueError("Duplicate Excel column headers.")
        return [{k: (str(v).strip() if v is not None else "") for k, v in zip(headers, r) if k} for r in rows[1:] if any(v is not None for v in r)]
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    headers = reader.fieldnames or []
    normalized = [h.strip() for h in headers]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Duplicate CSV column headers.")
    rows = []
    for row in reader:
        if None in row:
            raise ValueError("CSV row contains more values than its headers.")
        rows.append({str(k).strip(): str(v).strip() if v is not None else "" for k, v in row.items()})
    return rows


def required(rows, columns):
    if not rows:
        raise ValueError("No data rows.")
    missing = set(columns) - set(rows[0])
    if missing:
        raise ValueError(f"Missing fields: {', '.join(sorted(missing))}")


def samples_from_rows(rows, icp_rows):
    required(rows, ["sample_id", "origin", "target_elements"])
    if icp_rows:
        icp_rows = [normalize_assay(r) for r in icp_rows]
        required(icp_rows, ["value" if c == "w_wt_percent" else c for c in ICP_COLUMNS if c not in PRECURSOR_COLUMNS])
    ids = [r["sample_id"] for r in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate sample_id in the sample table.")
    if set(r["sample_id"] for r in icp_rows) - set(ids):
        raise ValueError("The ICP table contains sample IDs absent from the sample table.")
    grouped = defaultdict(list)
    for row in icp_rows:
        grouped[row["sample_id"]].append({k: row.get(k, "") for k in ["value" if c == "w_wt_percent" else c for c in ICP_COLUMNS] if k != "sample_id"})
    result = []
    for row in rows:
        elements = parse_elements(row["target_elements"])
        concentration = recipe(elements, finite_number(row["alpha"], "alpha", True)) if row.get("alpha") else {e: None for e in elements}
        assay = grouped[row["sample_id"]]
        for a in assay:
            if a.get("precursor_mM") and a["element"] in elements:
                concentration[a["element"]] = finite_number(a["precursor_mM"], "Element-specific concentration", True)
        sample = {k: row.get(k, "") for k in ("sample_id", "batch_id", "origin", "precursor_family", "support_batch", "protocol_id", "notes")}
        for key in ('pan_mass_g', 'solvent_volume_mL', 'preparation_basis', 'concentration_input_unit', 'electrical_input_kind'):
            if row.get(key) is not None and str(row.get(key)).strip():
                sample[key] = row[key]
        declared = {a['element']: precursor_from_row(a) for a in assay if a['element'] in elements and any(precursor_text(a.get(k)) for k in PRECURSOR_COLUMNS)}
        if declared:
            sample['precursors'] = declared
        sample["sample_condition"] = row.get("sample_condition") or "NOT_RECORDED"
        if row.get("record_uid"):
            sample["record_uid"] = row["record_uid"]
        if "qc" in row: sample["qc"] = row["qc"]  # Preserve legacy source values without using them for screening.
        sample.update(elements=elements, concentrations_mM=concentration, pulse_count=row.get("pulse_count", ""), temperature_C=row.get("temperature_C", ""), icp=assay,
                      icp_review={"state": row.get("icp_state") or "AUTO", "source": row.get("icp_source", ""), "reviewer": row.get("icp_reviewer", "")},
                      pulse_limit_C=row.get("pulse_limit_C", ""), voltage_V=row.get("voltage_V", ""), current_A=row.get("current_A", ""),
                      stem={"state": row.get("stem_state") or "NOT_TESTED", "source": row.get("stem_source", ""), "reviewer": row.get("stem_reviewer", ""), "notes": row.get("stem_notes", "")},
                      xrd={"state": row.get("xrd_state") or "PENDING", "mode": row.get("xrd_mode") or "AUTO", "source": row.get("xrd_source", ""), "criteria": row.get("xrd_criteria", ""), "reviewer": row.get("xrd_reviewer", "")})
        if row.get("recommendation_id"):
            sample["recommendation_id"] = row["recommendation_id"]
        if row.get("planned_temperature_C"):
            sample["planned_temperature_C"] = finite_number(row["planned_temperature_C"], "Planned temperature context", True)
        result.append(validate_sample(sample))
    return result


def methods_from_rows(rows):
    if rows and "loq_ug_g" in rows[0] and "loq_wt_percent" in rows[0]:
        raise ValueError("Use one set of method-limit columns, in wt%.")
    legacy = bool(rows and "loq_ug_g" in rows[0])
    columns = [c.replace("_wt_percent", "_ug_g") if legacy else c for c in METHOD_COLUMNS]
    required(rows, columns)
    grouped = {}
    for r in rows:
        ident = r["method_id"]
        if not ident:
            raise ValueError("method_id cannot be empty.")
        rule = r["rule"].upper()
        if rule not in {"LOD", "LOQ"} or r["origin"] not in {"REAL", "DEMO"}:
            raise ValueError("A method must specify LOD/LOQ and REAL/DEMO.")
        if r["confirmed"] not in {"0", "1"}:
            raise ValueError("confirmed must be 0 or 1.")
        if not r["version"] or not r["instrument"]:
            raise ValueError("Method version and instrument name cannot be empty.")
        if r["confirmed"] == "1" and not r["source"]:
            raise ValueError("A confirmed method requires a source for its limits.")
        obj = {"method_id": ident, "origin": r["origin"], "instrument": r["instrument"], "rule": rule, "confirmed": r["confirmed"] == "1", "source": r["source"], "version": r["version"], "limits": {}, "unit": "wt%"}
        if ident in grouped:
            prior = grouped[ident]
            if {k: v for k, v in prior.items() if k != "limits"} != {k: v for k, v in obj.items() if k != "limits"}:
                raise ValueError(f"{ident}'s method metadata are inconsistent.")
        else:
            grouped[ident] = obj
        if r["element"] not in ELEMENTS or r["element"] in grouped[ident]["limits"]:
            raise ValueError("Invalid or duplicate method element.")
        limit = {}
        for key in ("lod", "loq"):
            raw = r[f"{key}_ug_g" if legacy else f"{key}_wt_percent"]
            if raw and legacy:
                raw = to_wt_percent(raw, "ug/g")
            limit[key] = finite_number(raw, key, True) if raw else None
        if limit["lod"] and limit["loq"] and limit["loq"] < limit["lod"]:
            raise ValueError("LOQ cannot be lower than LOD.")
        if obj["confirmed"] and limit[rule.lower()] is None:
            raise ValueError(f"Confirmed {rule} criterion requires element-specific limits.")
        grouped[ident]["limits"][r["element"]] = limit
    return list(grouped.values())


def csv_bytes(rows, columns=None):
    out = io.StringIO()
    columns = columns or (list(rows[0]) if rows else [])
    writer = csv.DictWriter(out, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows([{**r, **({"w_wt_percent": r.get("w_wt_percent", r.get("value", ""))} if "w_wt_percent" in columns else {})} for r in rows])
    return out.getvalue().encode("utf-8-sig")


def sample_tables(samples):
    samples_rows, assay_rows = [], []
    for s in samples:
        row = {k: s.get(k, "") for k in SAMPLE_COLUMNS}
        row.update(target_elements=";".join(s["elements"]),
                   **{f"icp_{k}": s.get("icp_review", {}).get(k, "AUTO" if k == "state" else "") for k in ("state", "source", "reviewer")},
                   xrd_mode=s["xrd"].get("mode", "AUTO"),
                   **{f"stem_{k}": s.get("stem", {}).get(k, "NOT_TESTED" if k == "state" else "") for k in ("state", "source", "reviewer", "notes")},
                   xrd_state=s["xrd"].get("state", "PENDING"), xrd_source=s["xrd"].get("source", ""), xrd_criteria=s["xrd"].get("criteria", ""), xrd_reviewer=s["xrd"].get("reviewer", ""))
        samples_rows.append(row)
        for a in s.get("icp", []):
            assay_rows.append({**a, "sample_id": s["sample_id"], "precursor_mM": s["concentrations_mM"].get(a["element"], ""), **precursor_to_columns(s, a['element'])})
        known = {a["element"] for a in s.get("icp", [])}
        for e in s["elements"]:
            if e not in known:
                assay_rows.append({"sample_id": s["sample_id"], "element": e, "value": "", "unit": "wt%", "flag": "MISSING", "method_id": "", "source": "", "precursor_mM": s["concentrations_mM"][e], **precursor_to_columns(s, e)})
    return samples_rows, assay_rows


def recommendation_workbook(recommendations):
    """Create an unmeasured reporting workbook; never copy old outcomes."""
    import openpyxl
    sample_rows, icp_rows = [], []
    for r in recommendations:
        s = r["recipe"]
        row = {k: "" for k in SAMPLE_COLUMNS}
        row.update(sample_id=r["recommendation_id"], batch_id="REPLACE_WITH_SYNTHESIS_BATCH_ID", origin=r["origin"], target_elements=";".join(s["elements"]),
                   pulse_count=s["pulse_count"], temperature_C=s["temperature_C"], precursor_family=s["precursor_family"],
                   support_batch=s["support_batch"], protocol_id=s["protocol_id"], pulse_limit_C=s.get("pulse_limit_C", ""), xrd_state="PENDING", icp_state="AUTO", xrd_mode="MANUAL", stem_state="NOT_TESTED",
                   notes="Proposed recipe. Replace batch_id after actual synthesis; fill method IDs and measured results.")
        row.update({k: s[k] for k in ('pan_mass_g', 'solvent_volume_mL', 'preparation_basis') if k in s})
        sample_rows.append(row)
        for e in s["elements"]:
            icp_rows.append({"sample_id": r["recommendation_id"], "element": e, "value": "", "unit": "wt%", "flag": "MISSING", "method_id": "", "source": "", "precursor_mM": s["concentrations_mM"][e], **precursor_to_columns(s, e)})
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet, rows, columns in (("samples", sample_rows, SAMPLE_COLUMNS), ("icp", icp_rows, ICP_COLUMNS), ("methods", [], METHOD_COLUMNS)):
        ws = book.create_sheet(sheet)
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width = min(40, max(18, len(str(cells[0].value)) + 2))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
