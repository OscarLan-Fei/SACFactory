import os
import io
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from streamlit.testing.v1 import AppTest


class UITests(unittest.TestCase):
    def test_analysis_periodic_controls_apply_without_mutating_records_or_other_fields(self):
        import copy
        from sacfactory.storage import Store
        from sacfactory.domain import ELEMENTS
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        store = Store(self.tmp.name, 'DEMO')
        before = copy.deepcopy(store.all('sample'))
        controls = list(self.app.selectbox) + list(self.app.multiselect)
        count = 0
        for control in controls:
            if control.options and set(control.options).issubset(ELEMENTS):
                self.assertEqual(self.app.button(key=control.key + '_periodic_open').label,
                    'Periodic table: ' + control.label)
                count += 1
        self.assertGreaterEqual(count, 10)
        original_pairs = self.app.multiselect(key='pair_elements_DEMO').value
        self.app.button(key='weighted_elements_DEMO_periodic_open').click().run()
        self.app.button(key='periodic_draft_weighted_elements_DEMO_clear').click().run()
        self.app.button(key='periodic_draft_weighted_elements_DEMO_3_Fe').click().run()
        self.app.button(key='periodic_draft_weighted_elements_DEMO_3_Co').click().run()
        self.app.button(key='periodic_draft_weighted_elements_DEMO_apply').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.multiselect(key='weighted_elements_DEMO').value, ['Fe', 'Co'])
        self.assertEqual(self.app.multiselect(key='pair_elements_DEMO').value, original_pairs)
        self.assertEqual(store.all('sample'), before)

    def test_sigma_chlorides_registered_and_filled_for_new_batch_elements_in_both_spaces(self):
        import copy
        from sacfactory.storage import Store
        from sacfactory.domain import TARGETS
        for mode, origin in [('Real experiments · REAL', 'REAL'), ('Simulated demo · DEMO', 'DEMO')]:
            self.app.radio(key='mode').set_value(mode).run()
            store = Store(self.tmp.name, origin)
            before = copy.deepcopy(store.all('sample'))
            library = store.get('meta', 'reagent_library')['precursors']
            self.assertEqual(set(library), set(TARGETS))
            self.assertTrue(all(v['type'] == 'CHLORIDE' for v in library.values()))
            self.app.radio(key='page').set_value('Materials & consumables').run()
            defaults = next(w.value for w in self.app.dataframe
                if 'precursor_compound' in w.value.columns and len(w.value) == 50)
            self.assertEqual(defaults.set_index('element').loc['Au', 'precursor_compound'], 'AuCl3')
            self.app.radio(key='page').set_value('Samples & data').run()
            next(w for w in self.app.button if w.label == 'Clear selection').click().run()
            next(w for w in self.app.button if w.label == 'Co').click().run()
            details = next(w.value for w in self.app.dataframe
                if {'sample_id', 'precursor_compound'}.issubset(w.value.columns))
            self.assertEqual(set(details.element), {'Co'})
            self.assertTrue(all(details.precursor_compound == 'CoCl2\u00b76H2O'))
            self.assertEqual(len(self.app.exception), 0)
            self.assertEqual(store.all('sample'), before)

    def test_weighted_validation_module_masks_elements_and_exports_without_writes(self):
        import copy
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store=Store(self.tmp.name,'DEMO'); before=copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertIn('Weighted validation',[w.label for w in self.app.tabs])
        self.assertEqual(len(self.app.multiselect(key='weighted_elements_DEMO').options),118)
        self.app.multiselect(key='weighted_elements_DEMO').set_value(['Fe','Co','Ni']).run()
        self.app.selectbox(key='weighted_protocol_DEMO').set_value('Unseen element holdout').run()
        self.app.number_input(key='weighted_aux_members_DEMO').set_value(1)
        self.app.number_input(key='weighted_aux_epochs_DEMO').set_value(1)
        self.app.number_input(key='weighted_aux_draws_DEMO').set_value(4)
        self.app.button(key='weighted_run_DEMO').click().run(timeout=90)
        self.assertEqual(len(self.app.exception),0)
        self.assertGreater(len(self.app.session_state['weighted_result_DEMO'][1][0]),90)
        self.assertTrue(any(w.label=='Download prediction pairs' for w in self.app.get('download_button')))
        info=self.app.session_state['weighted_result_DEMO'][1][4]
        self.assertEqual(info['selected_elements'],['Fe','Co','Ni'])
        for fold in info['folds']:
            self.assertFalse(set(fold['train_batch_ids']) & set(fold['test_batch_ids']))
            heldout={s['elements'][0] for s in before if s['sample_id'] in fold['test_sample_ids']}
            self.assertFalse(heldout & set(fold['train_elements']))
        self.assertEqual(store.all('sample'),before)

    def test_combined_element_loading_ui_all50_selection_and_exports_preserve_records(self):
        import copy
        import json
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        before = copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Data analysis').run()
        next(w for w in self.app.button if w.label == 'Select all 50 campaign elements').click().run()
        selector = self.app.multiselect(key='combined_loading_elements_DEMO')
        self.assertEqual(len(selector.options), 118)
        self.assertEqual(len(selector.value), 50)
        # Selection has no 50-element cap; use an explicit four-element subset
        # to exercise chart fidelity without fitting fifty GP models in a UI test.
        self.app.multiselect(key='combined_loading_elements_DEMO').set_value(['Fe','Co','Ni','Cu']).run()
        self.app.selectbox(key='combined_loading_model_DEMO').set_value('Random forest').run()
        next(w for w in self.app.button if w.label == 'Evaluate and plot selected elements').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(next(w for w in self.app.metric if w.label == 'Plotted elements').value, '4')
        self.assertGreater(int(next(w for w in self.app.metric if w.label == 'Combined held-out samples').value), 70)
        self.assertTrue(any(w.label == 'Download combined loading points CSV' for w in self.app.get('download_button')))
        spec = next(json.loads(w.proto.spec) for w in self.app.get('vega_lite_chart')
                    if json.loads(w.proto.spec).get('layer', [{}])[0].get('encoding', {}).get('color', {}).get('field') == 'element')
        self.assertEqual(len(spec['layer'][0]['data']['values']), int(next(w for w in self.app.metric if w.label == 'Combined held-out samples').value))
        self.assertEqual(set(spec['layer'][0]['encoding']['color']['scale']['domain']), {'Fe', 'Co', 'Ni', 'Cu'})
        self.assertEqual(store.all('sample'), before)

    def test_element_campaign_ui_source_selection_loading_and_strategy(self):
        import copy
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        before = copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertIn('Element campaign', [t.label for t in self.app.tabs])
        self.assertEqual(next(w for w in self.app.metric if w.label == 'Single elements with screening candidates').value, '50 / 50')
        self.app.multiselect(key='element_sources_DEMO').set_value(['Fe']).run()
        self.app.multiselect(key='combined_loading_elements_DEMO').set_value(['Co']).run()
        self.app.button(key='combined_loading_run_DEMO').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Download combined loading points CSV' for w in self.app.get('download_button')))
        self.app.checkbox(key='element_campaign_policy_DEMO').check()
        next(w for w in self.app.button if w.label == 'Save campaign strategy').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(store.get('meta', 'element_campaign')['enabled'])
        self.assertEqual(store.all('sample'), before)

    def test_loading_ui_protected_predictions_and_missing_descriptor_explanations(self):
        import copy
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        before = copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Data analysis').run()
        self.app.radio(key='loading_data_mode_DEMO').set_value('Single-element model').run()
        self.assertIn('ICP loading prediction', [w.label for w in self.app.tabs])
        next(w for w in self.app.button if w.label == 'Validate ICP loading predictions').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertGreater(int(next(w for w in self.app.metric if w.label == 'Held-out loading samples').value), 8)
        self.assertTrue(any(w.label == 'Download held-out loading predictions CSV' for w in self.app.get('download_button')))
        variables = next(w for w in self.app.multiselect if w.label == 'Correlation variables')
        variables.set_value(['Fe_mM', 'formal_oxidation_state_mean', 'ionic_radius_pm_mean']).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any('formal_oxidation_state_mean: Constant values' in w.value for w in self.app.info))
        self.assertEqual(store.all('sample'), before)

    def test_prediction_points_toggle_preserves_eligibility_and_binary_outcomes(self):
        import copy
        import json
        from sacfactory.storage import Store
        from sacfactory.analysis_extended import prediction_records
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        before = copy.deepcopy(store.all('sample'))
        audit = prediction_records(before, {})
        valid_n = int(audit['eligible'].sum())
        self.app.radio(key='page').set_value('Data analysis').run()
        points = self.app.radio(key='prediction_plot_points_DEMO')
        self.assertEqual(points.value, 'Individual samples')
        self.assertEqual(int(next(w for w in self.app.metric if w.label == 'Auditable prediction/result pairs').value), valid_n)
        self.assertTrue(any(f'All {valid_n} eligible samples are plotted' in w.value for w in self.app.caption))
        spec = next(json.loads(w.proto.spec) for w in self.app.get('vega_lite_chart')
                    if json.loads(w.proto.spec).get('encoding', {}).get('x', {}).get('field') == 'p_screen_pass')
        self.assertEqual(spec['encoding']['y']['field'], 'label')
        self.assertEqual(spec['encoding']['y']['axis']['values'], [0, 1])
        for mode in ['Probability bins', 'Repeated recipes', 'Individual samples']:
            self.app.radio(key='prediction_plot_points_DEMO').set_value(mode).run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertEqual(int(next(w for w in self.app.metric if w.label == 'Auditable prediction/result pairs').value), valid_n)
        self.assertEqual(store.all('sample'), before)

    def test_overview_keeps_a_visible_correlation_entry_and_switches_scope(self):
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        self.app.radio(key='analysis_scope_DEMO').set_value('All cohorts: overview only').run()
        self.assertIn('Correlations & coverage', [t.label for t in self.app.tabs])
        self.assertIn('Variable coverage and missing values', [w.label for w in self.app.expander])
        next(w for w in self.app.button if w.label == 'Show cross-cohort correlations').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.radio(key='analysis_scope_DEMO').value, 'Cross-cohort exploratory associations')
        self.assertTrue(any(w.label == 'Correlation variables' for w in self.app.multiselect))
        self.assertTrue(any(w.label == 'Download correlations CSV' for w in self.app.get('download_button')))

    def test_interaction_chart_keeps_bin_steps_in_the_rendered_ui(self):
        import json
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        spec = next(json.loads(w.proto.spec) for w in self.app.get('vega_lite_chart')
                    if json.loads(w.proto.spec).get('encoding', {}).get('y', {}).get('field') == 'y_range')
        self.assertEqual(spec['height'], {'step': 44})
        self.assertEqual(spec['width'], {'step': 180})
        self.assertEqual(len(spec['encoding']['y']['scale']['domain']), 4)
        self.assertFalse(spec['encoding']['y']['axis']['labelOverlap'])

    def test_cross_composition_associations_and_chemical_outcomes_preserve_records(self):
        import copy
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        original = copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Data analysis').run()
        self.app.radio(key='analysis_scope_DEMO').set_value('Cross-cohort exploratory associations').run()
        self.assertEqual(len(self.app.exception),0)
        concentration_elements = next(w for w in self.app.multiselect if w.label == 'Concentration elements to include')
        self.assertEqual(len(concentration_elements.options),118)
        variables = next(w for w in self.app.multiselect if w.label == 'Correlation variables')
        for name in ['Fe_mM','Co_mM','Ni_mM','total_concentration_mM','element_count','en_pauling_mean','formal_oxidation_state_mean','label']:
            self.assertIn(name,variables.options)
        self.assertIn('Co_mM',variables.value)
        self.assertIn('Ni_mM',variables.value)
        variables.set_value(['Fe_mM','Co_mM','Ni_mM','total_concentration_mM','element_count','en_pauling_mean','covalent_radius_cordero_mean','atomic_radius_rahm_mean','pulse_count','label']).run()
        self.assertEqual(len(self.app.exception),0)
        self.assertEqual(len(next(w for w in self.app.multiselect if w.label == 'Correlation variables').value),10)
        self.assertIn('Success & failure',[t.label for t in self.app.tabs])
        metric = next(w for w in self.app.metric if w.label == 'Compared screening outcomes')
        self.assertEqual(int(metric.value),1474)
        self.assertTrue(any(w.label == 'Download chemical outcome comparison' for w in self.app.get('download_button')))
        next(w for w in self.app.selectbox if w.label == 'Outcome comparison metric').set_value('formal_oxidation_state_mean').run()
        self.assertEqual(len(self.app.exception),0)
        self.assertGreater(int(next(w for w in self.app.metric if w.label == 'Compared descriptor observations').value),0)
        self.assertEqual(store.all('sample'),original)

    def test_analysis_all_118_pair_elements_and_eight_correlation_axes(self):
        import json
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        selector = next(w for w in self.app.multiselect if w.label == 'Correlation variables')
        selector.set_value(selector.options[:8]).run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.button if w.label == 'Select all 118 elements').click().run()
        self.assertEqual(len(self.app.exception), 0)
        pairs = self.app.multiselect(key='pair_elements_DEMO')
        self.assertEqual(len(pairs.options), 118)
        self.assertEqual(len(pairs.value), 118)
        specs = [json.loads(w.proto.spec) for w in self.app.get('vega_lite_chart')]
        matrices = [spec for spec in specs if spec.get('encoding', {}).get('y', {}).get('axis', {}).get('labelOverlap') is False]
        pair_spec = next(spec for spec in matrices if spec['encoding']['y'].get('field') == 'element_2')
        self.assertEqual(len(pair_spec['encoding']['y']['scale']['domain']), 118)
        corr_spec = next(spec for spec in matrices if spec['encoding']['y'].get('field') == 'parameter_2')
        self.assertEqual(corr_spec['encoding']['y']['scale']['domain'], selector.options[:8])
        self.assertEqual(corr_spec['height'], {'step':32})

    def test_descriptor_audit_filters_and_planning_context_do_not_backfill_records(self):
        import copy
        from sacfactory.storage import Store
        from sacfactory.precursor_catalog import planning_draft
        store = Store(self.tmp.name, 'REAL')
        before = copy.deepcopy(store.all('sample'))
        context = copy.deepcopy(before[0])
        context['precursors'] = planning_draft(context['elements'])
        store.put_many('meta', [{'id':'plan_context', 'origin':'REAL', 'context':context}], 'id')
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertIn('Chemical descriptors', [w.label for w in self.app.tabs])
        metric = next(w for w in self.app.metric if w.label == 'Source-backed oxidation memberships')
        self.assertEqual(metric.value, '0')
        self.assertTrue(any(w.label == 'Download recipe chemistry summaries' for w in self.app.get('download_button')))
        self.assertEqual(store.all('sample'), before)

    def test_saved_reagent_defaults_follow_new_targets_and_preserve_original_samples(self):
        import copy
        from sacfactory.storage import Store
        from sacfactory.precursor_catalog import planning_draft
        store = Store(self.tmp.name, 'REAL')
        original = copy.deepcopy(store.all('sample'))
        library = planning_draft(['Cu', 'In', 'Au'])
        for entry in library.values():
            entry['source'] = 'User declared reagent; Sigma identity reference for synthetic UI fixture'
        store.put_many('meta', [{'id': 'reagent_library', 'origin': 'REAL', 'precursors': library}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertFalse(any('precursor_compound' in frame.value.columns for frame in self.app.dataframe))
        next(w for w in self.app.button if w.label == 'Clear selection').click().run()
        for element in ('Cu', 'In', 'Au'):
            next(w for w in self.app.button if w.label == element).click().run()
        next(w for w in self.app.checkbox if w.label == 'The entered operating range has been validated on the device').check()
        next(w for w in self.app.button if w.label == 'Save operating range').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        ctx = store.get('meta', 'plan_context')['context']
        self.assertEqual({e: v['compound'] for e, v in ctx['precursors'].items()}, {'Cu': 'CuCl2', 'In': 'InCl3', 'Au': 'AuCl3'})
        next(w for w in self.app.button if w.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(all(row['recipe']['precursors'] == ctx['precursors'] for row in self.app.session_state['preview_REAL']))
        self.assertEqual(store.all('sample'), original)
        self.app.radio(key='page').set_value('Materials & consumables').run()
        self.assertTrue(any(e.label == 'Saved reagent defaults' for e in self.app.expander))
        next(w for w in self.app.button if w.label == 'Load Sigma candidates into planning draft').click().run()
        from sacfactory.planning import fingerprint
        draft_key = f"catalog_draft_REAL_{fingerprint(ctx['elements'])}"
        self.assertTrue(all(not v['source'].startswith('PLANNED_CATALOG_ASSUMPTION:') for v in self.app.session_state[draft_key].values()))
        self.assertEqual(store.all('sample'), original)

    def test_material_setup_is_read_by_planning_without_rewriting_experiments(self):
        import copy
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        before = copy.deepcopy(store.all('sample'))
        self.app.radio(key='page').set_value('Materials & consumables').run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.text_input if w.label == 'Fixed precursor system / version').set_value('CHLORIDE_V2')
        next(w for w in self.app.text_input if w.label == 'Fixed support system / batch').set_value('CARBON_BATCH_2')

        def edited(frame, *args, **kwargs):
            frame = frame.copy()
            if 'precursor_compound' in frame.columns:
                frame.loc[frame.element == 'Fe', ['precursor_compound', 'precursor_type', 'precursor_oxidation_state', 'precursor_source']] = ['FeCl2', 'CHLORIDE', 2, 'Synthetic reagent record']
            return frame

        with patch('streamlit.data_editor', side_effect=edited):
            next(w for w in self.app.button if w.label == 'Save material setup').click().run()
        self.assertEqual(len(self.app.exception), 0)
        ctx = store.get('meta', 'plan_context')['context']
        self.assertEqual(ctx['precursors']['Fe']['compound'], 'FeCl2')
        self.assertEqual(store.all('sample'), before)
        next(w for w in self.app.button if w.label == 'Open Next batch').click().run()
        self.assertEqual(self.app.radio(key='page').value, 'Next batch')
        self.assertFalse(any(w.label == 'Fixed support system / batch' for w in self.app.text_input))
        next(w for w in self.app.checkbox if w.label == 'The entered operating range has been validated on the device').check()
        next(w for w in self.app.button if w.label == 'Save operating range').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        saved = store.get('meta', 'plan_context')['context']
        self.assertEqual(saved['support_batch'], 'CARBON_BATCH_2')
        self.assertEqual(saved['precursors']['Fe']['compound'], 'FeCl2')
        self.assertEqual(store.all('sample'), before)

    def test_inventory_partial_entry_namespace_isolation_and_row_deletion(self):
        import copy
        import pandas as pd
        from sacfactory.storage import Store
        from sacfactory.ui_materials import LEDGER_COLUMNS
        real = Store(self.tmp.name, 'REAL')
        before = copy.deepcopy(real.export())
        self.app.radio(key='page').set_value('Materials & consumables').run()

        def edited(frame, *args, **kwargs):
            if list(frame.columns) == LEDGER_COLUMNS:
                return pd.DataFrame([{'item_name': 'Carbon paper', 'category': 'Consumable', 'quantity': None, 'notes': 'Synthetic inventory fixture'}], columns=LEDGER_COLUMNS)
            return frame

        with patch('streamlit.data_editor', side_effect=edited):
            next(w for w in self.app.button if w.label == 'Save inventory').click().run()
        self.assertEqual(len(self.app.exception), 0)
        rows = real.get('meta', 'materials_consumables')['rows']
        self.assertEqual(rows[0]['item_name'], 'Carbon paper')
        self.assertIsNone(rows[0]['quantity'])
        after = real.export()
        for key in before:
            if key not in {'meta', 'audit'}:
                self.assertEqual(after[key], before[key])
        self.assertEqual(after['audit'][:-1], before['audit'])
        self.assertEqual(after['audit'][-1]['id'], 'materials_consumables')
        self.assertEqual(real.get('meta', 'plan_context'), None)
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.assertEqual(len(self.app.exception), 0)
        inventory = next(w.value for w in self.app.dataframe if list(w.value.columns) == LEDGER_COLUMNS)
        self.assertTrue(inventory.empty)
        self.app.radio(key='mode').set_value('Real experiments · REAL').run()

        def removed(frame, *args, **kwargs):
            return frame.iloc[0:0] if list(frame.columns) == LEDGER_COLUMNS else frame

        with patch('streamlit.data_editor', side_effect=removed):
            next(w for w in self.app.button if w.label == 'Save inventory').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(real.get('meta', 'materials_consumables')['rows'], [])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.previous_dir = os.environ.get("SACFACTORY_DATA_DIR")
        os.environ["SACFACTORY_DATA_DIR"] = self.tmp.name
        from sacfactory.storage import Store
        from sacfactory.domain import TARGETS, recipe
        from test_science import sample, method
        fixture = sample()
        fixture.update(sample_id="REAL_ICP_001", elements=list(TARGETS), concentrations_mM=recipe(TARGETS),
            icp=[{"element": e, "value": "", "unit": "wt%", "flag": "MISSING", "method_id": "M1", "source": ""} for e in TARGETS + ["Cd"]],
            xrd={"state": "PENDING"}, notes="Synthetic UI test fixture, not experimental data")
        limits = method()["M1"]
        limits["confirmed"] = False
        db = Store(self.tmp.name, "REAL")
        db.import_bundle([fixture], [limits])
        self.app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=30).run()

    def tearDown(self):
        if self.previous_dir is None: os.environ.pop("SACFACTORY_DATA_DIR", None)
        else: os.environ["SACFACTORY_DATA_DIR"] = self.previous_dir
        self.tmp.cleanup()

    def test_all_pages_real_and_demo_no_exceptions(self):
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.metric[0].value, "1")
        for mode in ("Real experiments · REAL", "Simulated demo · DEMO"):
            self.app.radio(key="mode").set_value(mode).run()
            for page in ("Experiment chat", "Samples & data", "Next batch", "Materials & consumables", "Data analysis", "Language model", "About"):
                self.app.radio(key="page").set_value(page).run()
                self.assertEqual(len(self.app.exception), 0, (mode, page, list(self.app.exception)))
        self.assertEqual(self.app.metric[0].value, "1476")

    def test_real_demo_shared_controls_chemistry_and_analysis_parity(self):
        import copy
        from sacfactory.storage import Store
        real = Store(self.tmp.name, 'REAL')
        demo = Store(self.tmp.name, 'DEMO')
        rows = copy.deepcopy(real.all('sample'))
        for row in rows:
            row['origin'] = 'DEMO'
        demo.import_bundle(rows, [])
        demo.put_many('meta', [{'id': 'seed-v1', 'origin': 'DEMO', 'source': 'Matched synthetic parity fixture'}], 'id', overwrite=True)
        controls = {}
        for mode, origin in [('Real experiments · REAL', 'REAL'), ('Simulated demo · DEMO', 'DEMO')]:
            self.app.radio(key='mode').set_value(mode).run()
            self.app.radio(key='page').set_value('Next batch').run()
            self.assertEqual(len(self.app.exception), 0)
            controls[origin] = {
                'selectboxes': [(w.label, tuple(w.options)) for w in self.app.selectbox],
                'inputs': [w.label for w in self.app.text_input],
                'numbers': [w.label for w in self.app.number_input],
                'buttons': [w.label for w in self.app.button],
                'tables': [tuple(w.value.columns) for w in self.app.dataframe],
            }
            self.assertFalse(any('precursor_compound' in columns for columns in controls[origin]['tables']))
            self.assertFalse(any(w.label == 'Fixed precursor system / version' for w in self.app.text_input))
            self.assertTrue(any(e.label == 'Change target elements' and not e.proto.expanded for e in self.app.expander))
            profiles = next(w for w in self.app.selectbox if w.label == 'Prediction descriptor profile')
            self.assertEqual(len(profiles.options), 8)
            self.app.radio(key='page').set_value('Materials & consumables').run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertTrue(any('precursor_compound' in frame.value.columns for frame in self.app.dataframe))
            self.assertTrue(any(w.label == 'Save inventory' for w in self.app.button))
            self.app.radio(key='page').set_value('Data analysis').run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertIn('Chemical Transfer', [t.label for t in self.app.tabs])
            self.assertTrue(any(w.label == 'Run chemical-transfer ablation' for w in self.app.button))
            next(w for w in self.app.button if w.label == 'Run chemical-transfer ablation').click().run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertTrue(any('No evaluable source/target split' in w.value for w in self.app.info))
        self.assertEqual(controls['REAL'], controls['DEMO'])

    def test_chemistry_ablation_ui_and_declared_precursor_pending_roundtrip(self):
        from test_chemistry import campaign
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        rows = campaign()
        store.put_many('sample', rows, 'sample_id')
        target = rows[-1]
        space = {'model_engine': 'reliability', 'descriptor_profile': 'chemistry_precursor', 'learning_iterations': 2,
                 'concentration_bounds_mM': {e: {'min_mM': .01, 'max_mM': .15} for e in target['elements']},
                 'pulses': [1, 2], 'pulse_limit_min_C': 900, 'pulse_limit_max_C': 1100, 'max_total_mM': .3,
                 'verified': True, 'baseline_repeat': False, 'repeat_strategy': 'none',
                 'range_source': 'Synthetic UI test bounds', 'candidate_budget': 256}
        store.put_many('meta', [{'id': 'plan_context', 'origin': 'REAL', 'context': target},
                               {'id': 'space', 'origin': 'REAL', 'space': space}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.number_input if w.label == 'Maximum experiments in the next batch').set_value(2).run()
        next(w for w in self.app.button if w.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        preview = self.app.session_state['preview_REAL']
        self.assertTrue(all(row['recipe']['precursors'] == target['precursors'] for row in preview))
        next(w for w in self.app.button if w.label == 'Add batch to pending').click().run()
        self.assertEqual(len(self.app.exception), 0)
        pending = [s for s in store.all('sample') if s.get('recommendation_id')]
        self.assertEqual(len(pending), 2)
        self.assertTrue(all(s['precursors'] == target['precursors'] for s in pending))
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.number_input if w.label == 'Chemistry learning iterations').set_value(2).run()
        next(w for w in self.app.button if w.label == 'Run chemical-transfer ablation').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Download chemical-transfer predictions' for w in self.app.get('download_button')))
        frame, info = self.app.session_state['chemistry_result_REAL'][1]
        from sacfactory.precursors import PROFILES
        self.assertEqual(set(frame.profile), set(PROFILES))
        self.assertEqual(info['ionic_rows_with_exact_state_match'], 0)

    def test_catalogue_stages_planning_only_and_changed_profile_clears_unreserved_preview(self):
        import copy
        from test_chemistry import campaign
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        rows = campaign()
        store.put_many('sample', rows, 'sample_id')
        target = rows[-1]
        space = {'model_engine': 'reliability', 'descriptor_profile': 'chemistry_ion', 'learning_iterations': 2,
                 'concentration_bounds_mM': {e: {'min_mM': .01, 'max_mM': .15} for e in target['elements']},
                 'pulses': [1, 2], 'pulse_limit_min_C': 900, 'pulse_limit_max_C': 1100, 'max_total_mM': .3,
                 'verified': True, 'baseline_repeat': False, 'repeat_strategy': 'none',
                 'range_source': 'Synthetic UI test bounds', 'candidate_budget': 256}
        store.put_many('meta', [{'id': 'plan_context', 'origin': 'REAL', 'context': target},
                               {'id': 'space', 'origin': 'REAL', 'space': space}], 'id', overwrite=True)
        original = copy.deepcopy(store.export())
        self.app.radio(key='page').set_value('Next batch').run()
        self.app.radio(key='page').set_value('Materials & consumables').run()
        next(w for w in self.app.button if w.label == 'Load Sigma candidates into planning draft').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(store.export(), original)
        chemistry = next(w.value for w in self.app.dataframe if 'precursor_compound' in w.value.columns)
        self.assertEqual(chemistry.loc[chemistry.element == 'Fe', 'precursor_compound'].iloc[0], 'FeCl2')
        self.assertTrue(chemistry.precursor_source.str.startswith('PLANNED_CATALOG_ASSUMPTION:').all())
        self.app.radio(key='page').set_value('Next batch').run()
        next(w for w in self.app.number_input if w.label == 'Maximum experiments in the next batch').set_value(2).run()
        next(w for w in self.app.button if w.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.session_state['preview_REAL']), 2)
        space['descriptor_profile'] = 'chemistry_state'
        store.put_many('meta', [{'id': 'space', 'origin': 'REAL', 'space': space}], 'id', overwrite=True)
        self.app.run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertNotIn('preview_REAL', self.app.session_state)
        self.assertTrue(any('Earlier unreserved preview cleared' in w.value for w in self.app.info))

    def test_sparse_demo_upgrade_settings_next_batch_prediction(self):
        from sacfactory.storage import Store
        from sacfactory.domain import recipe_key
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        original_samples = store.all('sample')
        store.put_many('meta', [{'id': 'space', 'origin': 'DEMO', 'space': {
            'model_engine': 'reliability', 'descriptor_profile': 'chemistry',
            'repeat_strategy': 'allocation', 'repeat_budget_fraction': .5,
            'repeat_goal_half_width': .15, 'repeat_cap': 20, 'learning_iterations': 8,
        }}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run(timeout=60)
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(next(w for w in self.app.text_input if w.label == 'Available pulse counts, comma-separated').value, '1,2,3,4,6,8')
        next(w for w in self.app.button if w.label == 'Generate batch preview').click().run(timeout=60)
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        preview = self.app.session_state['preview_DEMO']
        self.assertEqual(len(preview), 6)
        self.assertEqual(len({recipe_key(row['recipe']) for row in preview}), 6)
        self.assertTrue(all(row['origin'] == 'DEMO' and 0 <= row['p_screen_pass'] <= 1 for row in preview))
        self.assertTrue(all(row['recipe']['elements'] == ['Fe'] for row in preview))
        observed_fe=[s['concentrations_mM']['Fe'] for s in original_samples if s['elements']==['Fe']]
        self.assertTrue(all(min(observed_fe) <= row['recipe']['concentrations_mM']['Fe'] <= max(observed_fe) for row in preview))
        self.assertEqual(store.all('sample'), original_samples)
        self.assertEqual(store.all('recommendation'), [])

    def test_empty_demo_ranges_can_be_edited_and_preview_reports_validation(self):
        from sacfactory.storage import Store
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        store = Store(self.tmp.name, 'DEMO')
        store.put_many('meta', [{'id': 'space', 'origin': 'DEMO', 'space': {
            'concentration_bounds_mM': {}, 'model_engine': 'baseline',
        }}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.button if w.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any('concentration range' in row.value for row in self.app.error))

    def test_data_analysis_navigation_scope_filters_and_exports(self):
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        workspace = self.app.radio(key="page")
        self.assertNotIn("Analytical methods", workspace.options)
        self.assertEqual(workspace.options[workspace.options.index("Next batch") + 1], "Data analysis")
        workspace.set_value("Data analysis").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == "Download correlations CSV" for w in self.app.get("download_button")))
        next(w for w in self.app.radio if w.label == "Analysis scope").set_value("All cohorts: overview only").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(next(w for w in self.app.metric if w.label == "Selected records").value, "1476")
        self.assertFalse(any(w.label == "Parameter to summarize" for w in self.app.selectbox))
        next(w for w in self.app.radio if w.label == "Analysis scope").set_value("One comparable cohort").run()
        selector = next(w for w in self.app.selectbox if w.label == "Comparable composition and process")
        for index in range(len(selector.options)):
            self.app.selectbox(key="analysis_cohort_DEMO").set_value(index).run()
            self.assertEqual(len(self.app.exception), 0)
        batches = next(w for w in self.app.multiselect if w.label == "Filter by batch (empty = all)")
        batches.set_value(batches.options[:1]).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertGreater(int(next(w for w in self.app.metric if w.label == "Selected records").value), 0)

    def test_data_analysis_missing_parameters_do_not_crash(self):
        self.app.radio(key="page").set_value("Data analysis").run()
        parameter = next(w for w in self.app.selectbox if w.label == "Parameter to summarize")
        parameter.set_value("pulse_limit_C").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any("no recorded values" in item.value for item in self.app.info))
        self.assertEqual(next(w for w in self.app.metric if w.label == "Candidate pass rate").value, "Unavailable")

    def test_extended_analysis_demo_modules_surface_and_transfer(self):
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        labels = [t.label for t in self.app.tabs]
        for label in ('Outcome Maps', 'Repeatability', 'Composition & Transfer', 'Model Validation'):
            self.assertIn(label, labels)
        self.assertGreater(int(next(w for w in self.app.metric if w.label == 'Auditable prediction/result pairs').value), 80)
        next(w for w in self.app.button if w.label == 'Calculate conditional probability map').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Probability map display' for w in self.app.selectbox))
        for name in ('predictive_entropy_bits', 'nearest_training_distance'):
            next(w for w in self.app.selectbox if w.label == 'Probability map display').set_value(name).run()
            self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.button if w.label == 'Evaluate paired transfer benefit').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Transfer Brier improvement (positive = better)' for w in self.app.metric))
        next(w for w in self.app.multiselect if w.label == 'Model-guided synthesis batches').set_value(['DEMO_001']).run()
        next(w for w in self.app.multiselect if w.label == 'Comparator synthesis batches').set_value(['DEMO_002']).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Download strategy comparison CSV' for w in self.app.get('download_button')))

    def test_transfer_preview_one_click_pending_and_report_feedback(self):
        from test_transfer import source, context, space
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        store = Store(self.tmp.name, "REAL")
        store.put_many("sample", [source(i, passed=i >= 8) for i in range(18)], "sample_id")
        operating_space = {**space(), "element_transfer": True, "range_source": "Synthetic test range"}
        store.put_many("meta", [{"id": "plan_context", "origin": "REAL", "context": context()}, {"id": "space", "origin": "REAL", "space": operating_space}], "id", overwrite=True)
        self.app.radio(key="page").set_value("Next batch").run()
        self.assertEqual(len(self.app.exception), 0)
        next(w for w in self.app.number_input if w.label == "Maximum experiments in the next batch").set_value(2).run()
        next(b for b in self.app.button if b.label == "Generate batch preview").click().run()
        self.assertEqual(len(self.app.exception), 0)
        previews = self.app.session_state["preview_REAL"]
        self.assertTrue(all(r["source_elements"] == ["Fe"] for r in previews))
        self.assertTrue(all(r["recipe"]["elements"] == ["Co"] for r in previews))
        next(b for b in self.app.button if b.label == "Add batch to pending").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(screen(store.get("sample", "SAC_0001"), {})["status"], "PENDING")
        self.assertTrue(next(b for b in self.app.button if b.label == "Add batch to pending").disabled)
        self.app.radio(key="page").set_value("Samples & data").run()
        self.app.radio(key="record_filter_REAL").set_value("Pending samples").run()
        next(w for w in self.app.selectbox if w.label == "Inspect sample").set_value("SAC_0001").run()
        for label, value in (("Actual synthesis batch ID", "TEST_CO_ACTUAL_1"), ("Recorded measured peak / C", "1015")):
            next(w for w in self.app.text_input if w.label == label).set_value(value)
        next(w for w in self.app.selectbox if w.label == "ICP judgment").set_value("PASS")
        next(w for w in self.app.selectbox if w.label == "XRD relevant-phase criterion").set_value("PASS")
        next(b for b in self.app.button if b.label == "Save evidence revision").click().run()
        self.assertEqual(len(self.app.exception), 0)
        saved = store.get("sample", "SAC_0001")
        self.assertEqual(screen(saved, {})["status"], "PASS")
        self.assertEqual(store.get("recommendation", saved["recommendation_id"])["state"], "COMPLETED")

    def test_single_to_mixture_transfer_preview_and_pending_ui(self):
        from test_transfer import source
        from test_composition import mixture, mixture_space
        from sacfactory.storage import Store
        store = Store(self.tmp.name, "REAL")
        store.put_many("sample", [source(i, "Fe" if i < 18 else "Co", passed=i % 18 >= 8) for i in range(36)], "sample_id")
        operating = {**mixture_space(), "element_transfer": True, "range_source": "Synthetic mixture test range"}
        store.put_many("meta", [{"id": "plan_context", "origin": "REAL", "context": mixture()}, {"id": "space", "origin": "REAL", "space": operating}], "id", overwrite=True)
        self.app.radio(key="page").set_value("Next batch").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == "Use existing single-element and mixture data for transfer" for w in self.app.checkbox))
        next(w for w in self.app.number_input if w.label == "Maximum experiments in the next batch").set_value(2).run()
        next(b for b in self.app.button if b.label == "Generate batch preview").click().run()
        self.assertEqual(len(self.app.exception), 0)
        previews = self.app.session_state["preview_REAL"]
        self.assertTrue(all(r["unobserved_pair_count"] == 1 for r in previews))
        self.assertTrue(all(set(r["recipe"]["elements"]) == {"Fe", "Co"} for r in previews))
        next(b for b in self.app.button if b.label == "Add batch to pending").click().run()
        self.assertEqual(len(self.app.exception), 0)
        saved = store.get("sample", "SAC_0001")
        self.assertEqual(set(saved["elements"]), {"Fe", "Co"})
        self.assertEqual(saved["proposal"]["model_info"]["source_stages"], {"1": 36})

    def test_real_dialogue_diagnosis_and_next_batch(self):
        self.app.chat_input[0].set_value("Diagnose REAL_ICP_001").run()
        self.assertEqual(len(self.app.exception), 0)
        text = "\n".join(x.value for x in self.app.markdown)
        self.assertIn("Evidence pending", text)
        self.app.chat_input[0].set_value("Recommend 6 experiments").run()
        self.assertEqual(len(self.app.exception), 0)
        previews = self.app.session_state["preview_REAL"]
        self.assertEqual(len(previews), 1)
        self.assertIsNone(previews[0]["p_screen_pass"])

    def test_demo_recommendation_and_reservation(self):
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Next batch").run()
        next(b for b in self.app.button if b.label == "Generate batch preview").click().run()
        self.assertEqual(len(self.app.exception), 0)
        previews = self.app.session_state["preview_DEMO"]
        self.assertGreater(len(previews), 0)
        next(b for b in self.app.button if b.label == "Reserve batch").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(e.label == "Reserved work queue" for e in self.app.expander))

    def test_navigation_shortcut(self):
        next(b for b in self.app.button if b.label == "Configure composition and range").click().run()
        self.assertEqual(self.app.radio(key="page").value, "Next batch")
        self.assertEqual(len(self.app.exception), 0)

    def test_clear_chat_history_preserves_experiments_and_removes_starter_dialogue(self):
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        before = store.all('sample')
        self.app.chat_input[0].set_value('Summarize results').run()
        self.assertGreater(len(store.messages()), 0)
        next(b for b in self.app.button if b.label == 'Clear chat history').click().run()
        self.assertEqual(store.messages(), [])
        self.assertEqual(store.all('sample'), before)
        self.assertEqual(len(self.app.chat_message), 0)
        self.assertEqual(len(self.app.exception), 0)

    def test_adaptive_engine_policy_switch_and_prediction_metadata(self):
        from test_transfer import source, context, space
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        store.put_many('sample', [source(i, passed=i >= 8) for i in range(18)], 'sample_id')
        operating = {**space(), 'element_transfer': True, 'learning_iterations': 4,
                     'range_source': 'Synthetic adaptive UI test range'}
        store.put_many('meta', [{'id': 'plan_context', 'origin': 'REAL', 'context': context()},
                                {'id': 'space', 'origin': 'REAL', 'space': operating}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.selectbox(key='prediction_engine_REAL').value, 'adaptive')
        self.app.selectbox(key='acquisition_policy_REAL').set_value('exploration').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(store.get('meta', 'space')['space']['acquisition_policy'], 'exploration')
        next(b for b in self.app.button if b.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        rows = self.app.session_state['preview_REAL']
        self.assertTrue(all(r['policy'] == 'ADAPTIVE_JOINT_INFORMATION' for r in rows))
        self.assertTrue(all(r['acquisition']['name'] == 'exploration' for r in rows))
        next(b for b in self.app.button if b.label == 'Add batch to pending').click().run()
        pending = next(s for s in store.all('sample') if s.get('recommendation_id'))
        self.assertIn('p_screen_low', pending['proposal'])
        self.app.selectbox(key='prediction_engine_REAL').set_value('baseline').run()
        self.assertEqual(store.get('meta', 'space')['space']['model_engine'], 'baseline')
        self.assertEqual(len(self.app.exception), 0)

    def test_advanced_validation_ui_replay_and_strict_element(self):
        from test_research_validation import campaign
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        store.put_many('sample', campaign('Fe', 'F') + campaign('Co', 'C'), 'sample_id')
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        self.app.multiselect(key='advanced_targets_REAL').set_value([1])
        self.app.number_input(key='advanced_iterations_REAL').set_value(2)
        next(b for b in self.app.button if b.label == 'Run advanced validation').click().run()
        self.assertEqual(len(self.app.exception), 0)
        frame, info = self.app.session_state['advanced_result_REAL'][1]
        self.assertFalse(frame.empty)
        self.assertEqual(info['protocol'], 'CHRONOLOGICAL_CORRECTED_DATA_REPLAY')
        self.app.selectbox(key='advanced_protocol_REAL').set_value('Strict unseen element')
        self.app.multiselect(key='advanced_elements_REAL').set_value(['Co'])
        next(b for b in self.app.button if b.label == 'Run advanced validation').click().run()
        self.assertEqual(len(self.app.exception), 0)
        frame, info = self.app.session_state['advanced_result_REAL'][1]
        self.assertEqual(len(frame), 36)
        self.assertNotIn('Co', info['folds'][0]['train_elements'])

    def test_agreement_figure_controls_and_vector_download(self):
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Download agreement figure and source data' for w in self.app.get('download_button')))
        self.app.checkbox(key='agreement_stage_DEMO').set_value(False)
        self.app.checkbox(key='agreement_fit_DEMO').set_value(True).run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any(w.label == 'Download agreement figure and source data' for w in self.app.get('download_button')))
        self.app.radio(key='agreement_mode_DEMO').set_value('Repeated recipes').run()
        self.assertEqual(len(self.app.exception), 0)

    def test_critical_repeat_controls_persist_and_pending_preserves_rationale(self):
        from test_research_validation import campaign
        from test_transfer import context, space
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        store.put_many('sample', campaign(), 'sample_id')
        operating = {**space('Fe'), 'baseline_repeat': True, 'repeat_strategy': 'critical',
                     'batch_effect': True, 'learning_iterations': 2, 'range_source': 'Synthetic test bounds'}
        store.put_many('meta', [{'id': 'plan_context', 'origin': 'REAL', 'context': context('Fe')},
                                {'id': 'space', 'origin': 'REAL', 'space': operating}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(next(w for w in self.app.selectbox if w.label == 'Repeat selection').value, 'critical')
        next(b for b in self.app.button if b.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        repeat = next(r for r in self.app.session_state['preview_REAL'] if r.get('is_repeat'))
        next(b for b in self.app.button if b.label == 'Add batch to pending').click().run()
        saved = next(s for s in store.all('sample') if s.get('recommendation_id') == repeat['recommendation_id'])
        self.assertEqual(saved['proposal']['repeat_strategy'], 'critical')
        self.assertEqual(saved['proposal']['repeat_rationale'], repeat['repeat_rationale'])
        self.assertEqual(saved['batch_id'], 'UNKNOWN')

    def test_reliability_engine_allocation_preview_pending_and_recipe_analysis(self):
        from test_research_validation import campaign
        from test_transfer import context, space
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        store.put_many('sample', campaign()+campaign('Co', 'C'), 'sample_id')
        operating = {**space('Fe'), 'model_engine': 'reliability', 'descriptor_profile': 'chemistry',
            'baseline_repeat': True, 'repeat_strategy': 'allocation', 'repeat_budget_fraction': .5,
            'repeat_goal_half_width': .1, 'repeat_cap': 20, 'learning_iterations': 2,
            'range_source': 'Synthetic reliability test bounds', 'element_transfer': True}
        store.put_many('meta', [{'id': 'plan_context', 'origin': 'REAL', 'context': context('Fe')},
                                {'id': 'space', 'origin': 'REAL', 'space': operating}], 'id', overwrite=True)
        self.app.radio(key='page').set_value('Next batch').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.selectbox(key='prediction_engine_REAL').value, 'reliability')
        next(b for b in self.app.button if b.label == 'Generate batch preview').click().run()
        self.assertEqual(len(self.app.exception), 0)
        rows = self.app.session_state['preview_REAL']
        self.assertEqual(len(rows), 6)
        self.assertTrue(any(r.get('repeat_rationale', {}).get('planned_additional_repeats') for r in rows if r.get('repeat_rationale')))
        next(b for b in self.app.button if b.label == 'Add batch to pending').click().run()
        self.assertEqual(len(self.app.exception), 0)
        repeated = next(s for s in store.all('sample') if (s.get('proposal') or {}).get('repeat_strategy') == 'allocation')
        self.assertIn('further_repeats_after_this', repeated['proposal']['repeat_rationale'])
        self.app.radio(key='page').set_value('Data analysis').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertIn('Recipe Reliability', [t.label for t in self.app.tabs])

    def test_reliability_model_benchmark_ui_has_comparison_and_provenance_exports(self):
        from test_research_validation import campaign
        from sacfactory.storage import Store
        store = Store(self.tmp.name, 'REAL')
        store.put_many('sample', campaign()+campaign('Co', 'C'), 'sample_id')
        self.app.radio(key='page').set_value('Data analysis').run()
        self.app.multiselect(key='advanced_targets_REAL').set_value([1])
        self.app.selectbox(key='advanced_protocol_REAL').set_value('Reliability model benchmark').run()
        self.app.number_input(key='advanced_iterations_REAL').set_value(2)
        self.app.number_input(key='advanced_folds_REAL').set_value(2)
        next(b for b in self.app.button if b.label == 'Run advanced validation').click().run()
        self.assertEqual(len(self.app.exception), 0)
        frame, info = self.app.session_state['advanced_result_REAL'][1]
        self.assertEqual(frame.model.nunique(), 6)
        self.assertEqual(len(info['comparison']), 6)
        self.assertTrue(any(w.label == 'Download validation provenance' for w in self.app.get('download_button')))

    def test_bulk_upload_preview_commit_and_uploader_reset(self):
        import openpyxl
        root = Path(__file__).resolve().parents[1]
        book = openpyxl.load_workbook(root / "examples/bulk_import_template.xlsx")
        for name in ("samples", "methods"):
            ws = book[name]
            col = next(c.column for c in ws[1] if c.value == "origin")
            for row in range(2, ws.max_row + 1):
                ws.cell(row, col, "DEMO")
        upload = io.BytesIO()
        book.save(upload)
        upload.name = "pending_demo_batch.xlsx"

        def uploaded(label, *args, **kwargs):
            if label.startswith("Standard Excel"):
                return [upload] if kwargs["key"].endswith("_0") else []
            return None

        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Samples & data").run()
        with patch("streamlit.file_uploader", side_effect=uploaded):
            self.app.run()
            self.assertEqual(len(self.app.exception), 0)
            next(b for b in self.app.button if b.label == "Commit import to local database").click().run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertEqual(self.app.metric[0].value, "1477")
            self.assertEqual(self.app.session_state["upload_nonce_DEMO"], 1)
            self.assertEqual(len(self.app.error), 0)
            self.assertTrue(any("1 samples" in item.value for item in self.app.success))

    def test_full_experiment_record_edit_saves_conditions_content_and_judgments(self):
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.multiselect if w.label == "Record target elements").set_value(["Fe"]).run()
        for label, value in (("Record actual batch ID", "MANUAL_REAL_BATCH"), ("Record pulse count", "3"), ("Record controlled cutoff / C", "1050"),
                             ("Record measured peak / C", "1071"), ("Record voltage / V", "12"), ("Record current / A", "4")):
            next(w for w in self.app.text_input if w.label == label).set_value(value)
        next(w for w in self.app.selectbox if w.label == "Record ICP judgment").set_value("PASS")
        next(w for w in self.app.selectbox if w.label == "Record XRD judgment").set_value("PASS")

        def edit(frame, *args, **kwargs):
            frame = frame.copy()
            if kwargs.get("key", "").startswith("manual_records_") and "value" in frame:
                frame.loc[frame.element == "Fe", ["precursor_mM", "value"]] = [.07, "0.25"]
                frame.loc[frame.element == 'Fe', ['precursor_compound', 'precursor_type',
                    'precursor_oxidation_state', 'precursor_source']] = ['FeCl2', 'CHLORIDE', 2, 'Synthetic declared reagent']
            return frame

        with patch("streamlit.data_editor", side_effect=edit):
            next(b for b in self.app.button if b.label == "Save experiment changes").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        store = Store(self.tmp.name, "REAL")
        saved = store.get("sample", "REAL_ICP_001")
        self.assertEqual(saved["elements"], ["Fe"])
        self.assertEqual(saved["concentrations_mM"], {"Fe": .07})
        self.assertEqual(saved["icp"][0]["value"], "0.25")
        self.assertEqual(saved["icp"][0]["unit"], "wt%")
        self.assertEqual(saved["pulse_count"], 3)
        self.assertEqual(saved["pulse_limit_C"], 1050)
        self.assertEqual(saved["temperature_C"], 1071)
        self.assertEqual(saved['precursors']['Fe']['compound'], 'FeCl2')
        self.assertEqual(saved['precursors']['Fe']['source'], 'Synthetic declared reagent')
        self.assertEqual(screen(saved, {})["status"], "PASS")
        self.assertTrue(any(a["action"] == "manual_records" for a in store.export()["audit"]))

    def test_pending_editor_amends_plan_without_adding_observations(self):
        from sacfactory.storage import Store
        from sacfactory.planning import fit_model, recommend, fingerprint
        from sacfactory.domain import RULE_VERSION
        from test_transfer import context, space
        store = Store(self.tmp.name, "REAL")
        ctx, operating = context(), space()
        samples, methods = store.all("sample"), {m["method_id"]: m for m in store.all("method")}
        rows, _ = recommend(fit_model(samples, methods, ctx, transfer=True), samples, ctx, operating, "REAL", 1)
        rows[0]["snapshot"] = fingerprint({"samples": samples, "methods": methods, "context": ctx, "space": operating, "rule_version": RULE_VERSION})
        ident = store.enqueue_pending(rows, ctx, operating)[0]
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.selectbox if w.label == "Pending sample to edit").set_value(ident).run()
        next(w for w in self.app.text_input if w.label == "Planned pulse count").set_value("4")
        next(w for w in self.app.text_input if w.label == "Planned cutoff / C").set_value("1080")
        def edit_precursor(frame, *args, **kwargs):
            frame = frame.copy()
            if kwargs.get('key', '').startswith('manual_pending_') and 'precursor_compound' in frame:
                frame['precursor_compound'] = frame.element + 'Cl2'
                frame['precursor_source'] = 'Synthetic planned precursor revision'
            return frame
        with patch('streamlit.data_editor', side_effect=edit_precursor):
            next(b for b in self.app.button if b.label == "Save pending changes").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        saved = store.get("sample", ident)
        self.assertEqual(saved["pulse_count"], 4)
        self.assertEqual(saved["pulse_limit_C"], 1080)
        self.assertIsNone(saved["temperature_C"])
        self.assertEqual(saved["icp"], [])
        self.assertEqual(saved['precursors'][ctx['elements'][0]]['source'], 'Synthetic planned precursor revision')
        self.assertEqual(store.get('recommendation', rows[0]['recommendation_id'])['recipe']['precursors'], saved['precursors'])
        self.assertTrue(store.get("recommendation", rows[0]["recommendation_id"])["prediction_stale"])

    def test_manual_batch_feedback_completes_linked_samples_in_one_save(self):
        from sacfactory.storage import Store
        from sacfactory.planning import fit_model, recommend, fingerprint
        from sacfactory.domain import RULE_VERSION, screen
        from test_transfer import context, space
        store = Store(self.tmp.name, "REAL")
        ctx, operating = context(), space()
        samples, methods = store.all("sample"), {m["method_id"]: m for m in store.all("method")}
        rows, _ = recommend(fit_model(samples, methods, ctx, transfer=True), samples, ctx, operating, "REAL", 2)
        for row in rows:
            row["snapshot"] = fingerprint({"samples": samples, "methods": methods, "context": ctx, "space": operating, "rule_version": RULE_VERSION})
        ids = store.enqueue_pending(rows, ctx, operating)
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.multiselect if w.label == "Samples to return").set_value(ids).run()
        next(w for w in self.app.text_input if w.label == "Actual batch ID for this return (optional)").set_value("MANUAL_RETURN_BATCH")
        for label, value in (("Shared precursor system (optional)", "TEST_PRECURSOR"), ("Shared support system (optional)", "TEST_SUPPORT"), ("Shared protocol (optional)", "TEST_PROTOCOL")):
            next(w for w in self.app.text_input if w.label == label).set_value(value)

        def feedback(frame, *args, **kwargs):
            frame = frame.copy()
            key = kwargs.get("key", "")
            if key.startswith("feedback_outcomes_"):
                frame["icp_state"] = "PASS"
                frame["xrd_state"] = "PASS"
                frame["temperature_C"] = 1021
            if key.startswith("feedback_icp_"):
                frame["value"] = "0.25"
            return frame

        with patch("streamlit.data_editor", side_effect=feedback):
            next(b for b in self.app.button if b.label == "Save returned results").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        for ident in ids:
            sample = store.get("sample", ident)
            self.assertEqual(sample["batch_id"], "MANUAL_RETURN_BATCH")
            self.assertEqual(sample["icp"][0]["value"], "0.25")
            self.assertEqual(sample["temperature_C"], 1021)
            self.assertEqual(sample["precursor_family"], "TEST_PRECURSOR")
            self.assertEqual(sample["support_batch"], "TEST_SUPPORT")
            self.assertEqual(sample["protocol_id"], "TEST_PROTOCOL")
            self.assertEqual(screen(sample, methods)["status"], "PASS")
            self.assertEqual(store.get("recommendation", sample["recommendation_id"])["state"], "COMPLETED")
        self.assertTrue(any("2 completed" in item.value for item in self.app.success))

    def configure_api(self):
        self.app.radio(key="page").set_value("Language model").run()
        # Endpoint editing is reactive and clears the previous password widget.
        self.app.text_input(key='api_base_url').set_value('https://api.example.test/v1').run()
        for label, value in (("Model ID", "test-model"), ("API key", "test-private-key")):
            next(w for w in self.app.text_input if w.label == label).set_value(value)
        next(b for b in self.app.button if b.label == "Apply session settings").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertFalse(self.app.session_state["api_permission_REAL"])
        next(w for w in self.app.checkbox if w.label.startswith("Enable API")).check().run()
        self.assertTrue(self.app.session_state["api_permission_REAL"])

    def test_record_material_condition_can_fail_without_analytical_results(self):
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.selectbox if w.label == "Record sample condition").set_value("BURNED_OUT")
        next(w for w in self.app.text_input if w.label == "Record actual batch ID").set_value("ACTUAL_BURNOUT_BATCH")
        next(b for b in self.app.button if b.label == "Save experiment changes").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        saved = Store(self.tmp.name, "REAL").get("sample", "REAL_ICP_001")
        self.assertEqual(saved["sample_condition"], "BURNED_OUT")
        self.assertTrue(all(not a["value"] for a in saved["icp"]))
        self.assertEqual(screen(saved, {})["label"], 0)
        self.assertEqual(self.app.metric[2].value, "1")

    def test_delete_ui_compacts_sac_ids_and_keeps_manual_names(self):
        from sacfactory.storage import Store
        from test_science import sample
        store = Store(self.tmp.name, "REAL")
        store.put_many("sample", [{**sample(i), "sample_id": f"SAC_{i:04d}", "notes": f"Physical sample {i}"} for i in range(1, 4)], "sample_id")
        self.app.radio(key="page").set_value("Samples & data").run()
        self.assertTrue(next(w for w in self.app.checkbox if w.label == "Realign SAC IDs after deletion").value)
        next(w for w in self.app.multiselect if w.label == "Samples to delete").set_value(["SAC_0002"]).run()
        next(b for b in self.app.button if b.label == "Delete selected samples").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(store.get("sample", "SAC_0002")["notes"], "Physical sample 3")
        self.assertIsNotNone(store.get("sample", "REAL_ICP_001"))
        self.assertEqual(store.next_sample_ids(), ["SAC_0003"])
        self.assertTrue(any("next ID: SAC_0003" in message.value for message in self.app.success))

    def test_feedback_burnout_completes_pending_without_fake_fail_measurements(self):
        from sacfactory.storage import Store
        from sacfactory.planning import fit_model, recommend, fingerprint
        from sacfactory.domain import RULE_VERSION, screen
        from test_transfer import context, space
        store = Store(self.tmp.name, "REAL")
        ctx, operating = context(), space()
        samples, methods = store.all("sample"), {m["method_id"]: m for m in store.all("method")}
        proposals, _ = recommend(fit_model(samples, methods, ctx, transfer=True), samples, ctx, operating, "REAL", 1)
        proposals[0]["snapshot"] = fingerprint({"samples": samples, "methods": methods, "context": ctx, "space": operating, "rule_version": RULE_VERSION})
        ident = store.enqueue_pending(proposals, ctx, operating)[0]
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.multiselect if w.label == "Samples to return").set_value([ident]).run()
        next(w for w in self.app.text_input if w.label == "Actual batch ID for this return (optional)").set_value("ACTUAL_BURNOUT_RETURN")

        def burned(frame, *args, **kwargs):
            frame = frame.copy()
            if kwargs.get("key", "").startswith("feedback_outcomes_"):
                frame["sample_condition"] = "BURNED_OUT"
            return frame

        with patch("streamlit.data_editor", side_effect=burned):
            next(b for b in self.app.button if b.label == "Save returned results").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        saved = store.get("sample", ident)
        result = screen(saved, methods)
        self.assertEqual((result["status"], result["icp"], result["xrd"]), ("BURNED_OUT", "PENDING", "PENDING"))
        self.assertEqual(store.get("recommendation", proposals[0]["recommendation_id"])["state"], "COMPLETED")

    def test_api_settings_typed_request_local_shortcuts_and_namespace_permission(self):
        from sacfactory.storage import Store
        self.configure_api()
        self.app.radio(key="page").set_value("Experiment chat").run()
        parsed = {"intent": "recommend", "sample_id": None, "batch_size": 2, "pulse_count_filter": [1], "topic": None}
        with patch("sacfactory.language.CompatibleAdapter.interpret", return_value=parsed) as mock:
            self.app.chat_input[0].set_value("\u5e2e\u6211\u5b89\u6392\u4e24\u4e2a\u5b9e\u9a8c").run()
            self.assertEqual(len(self.app.exception), 0)
            mock.assert_called_once()
            self.assertEqual(mock.call_args.args[2]["origin"], "REAL")
            self.assertEqual(len(self.app.session_state["preview_REAL"]), 1)
            self.assertIsNone(self.app.session_state["preview_REAL"][0]["p_screen_pass"])
        store = Store(self.tmp.name, "REAL")
        self.assertNotIn("test-private-key", str(store.export()))
        self.assertEqual(store.messages()[-1]["payload"]["language"]["intent"], parsed)
        self.assertEqual(store.all("recommendation"), [])
        with patch("sacfactory.language.CompatibleAdapter.interpret") as mock:
            next(b for b in self.app.button if b.label == "Summary").click().run()
            mock.assert_not_called()
            self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
            self.app.chat_input[0].set_value("Summarize results").run()
            mock.assert_not_called()
        self.assertEqual(len(self.app.exception), 0)

    def test_openai_research_preset_sharing_multiturn_and_local_shortcuts(self):
        from sacfactory.storage import Store
        from sacfactory.dialogue import Answer
        from sacfactory.research_chat import ResearchAnswer
        self.app.radio(key='page').set_value('Language model').run()
        next(b for b in self.app.button if b.label == 'Use OpenAI preset').click().run()
        self.assertEqual(next(w for w in self.app.text_input if w.label == 'API base URL').value, 'https://api.openai.com/v1')
        next(w for w in self.app.text_input if w.label == 'Model ID').set_value('test-research-model')
        next(w for w in self.app.text_input if w.label == 'API key').set_value('test-private-key')
        next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.session_state['api_config'].api_format, 'responses')
        next(w for w in self.app.multiselect if w.label == 'Shared experiment records').set_value(['REAL_ICP_001']).run()
        next(w for w in self.app.checkbox if w.label.startswith('Enable API')).check().run()
        self.app.radio(key='page').set_value('Experiment chat').run()
        with patch('sacfactory.research_chat.ResearchAdapter.converse') as mock:
            self.app.chat_input[0].set_value('Discuss the synthesis results').run()
            mock.assert_not_called()
        self.app.radio(key='page').set_value('Language model').run()
        next(w for w in self.app.checkbox if w.label.startswith('Share selected experimental records')).check().run()
        self.app.radio(key='page').set_value('Experiment chat').run()
        reply = ResearchAnswer(Answer('These are recorded screening candidates; further tests can distinguish hypotheses.', [], []), [], {'input_tokens': 100, 'output_tokens': 20}, 1)
        with patch('sacfactory.research_chat.ResearchAdapter.converse', return_value=reply) as mock:
            self.app.chat_input[0].set_value('Discuss the Fe results').run()
            self.assertEqual(len(self.app.exception), 0)
            self.assertEqual(mock.call_args.args[2]['shared_records'][0]['sample_id'], 'REAL_ICP_001')
            self.app.chat_input[0].set_value('Explain the same hypothesis further').run()
            self.assertEqual(len(mock.call_args.args[1]), 2)
            next(b for b in self.app.button if b.label == 'Start new research discussion').click().run()
            self.app.chat_input[0].set_value('Begin a fresh discussion').run()
            self.assertEqual(mock.call_args.args[1], [])
            calls = mock.call_count
            next(b for b in self.app.button if b.label == 'Summary').click().run()
            self.assertEqual(mock.call_count, calls)
        store = Store(self.tmp.name, 'REAL')
        self.assertNotIn('test-private-key', str(store.export()))
        self.assertEqual(store.all('recommendation'), [])
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        with patch('sacfactory.research_chat.ResearchAdapter.converse') as mock:
            self.app.chat_input[0].set_value('Discuss simulated data').run()
            mock.assert_not_called()

    def test_parley_missing_model_retains_key_and_catalog_selection_applies(self):
        from sacfactory.storage import Store
        from sacfactory.providers import MIT_PARLEY_BASE
        self.app.radio(key='page').set_value('Language model').run()
        next(b for b in self.app.button if b.label == 'Use MIT Parley preset').click().run()
        self.assertEqual(self.app.text_input(key='api_base_url').value, MIT_PARLEY_BASE)
        self.app.text_input(key='api_key_input').set_value('sk-parley-v1-test-only')
        with patch('sacfactory.providers.post_json') as network:
            next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
            network.assert_not_called()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any('Model ID is missing' in e.value for e in self.app.error))
        self.assertNotIn('api_config', self.app.session_state)
        self.assertEqual(self.app.session_state['api_pending_config'].api_key, 'sk-parley-v1-test-only')
        self.app.text_input(key='api_key_input').set_value('')
        with patch('sacfactory.providers.post_json', return_value={'data': [{'id': 'school-model-a'}, {'id': 'school-model-b'}]}) as network:
            next(b for b in self.app.button if b.label == 'Load available models').click().run()
            network.assert_called_once()
            self.assertEqual(network.call_args.args[0].get_method(), 'GET')
            self.assertIsNone(network.call_args.args[0].data)
        self.assertEqual(len(self.app.exception), 0)
        self.app.selectbox(key='api_catalog_choice').set_value('school-model-b').run()
        self.assertEqual(self.app.text_input(key='api_model_id').value, 'school-model-b')
        self.app.text_input(key='api_key_input').set_value('')
        with patch('sacfactory.providers.post_json') as network:
            next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
            network.assert_not_called()
        self.assertEqual(len(self.app.exception), 0)
        config = self.app.session_state['api_config']
        self.assertEqual((config.base_url, config.model, config.api_format), (MIT_PARLEY_BASE, 'school-model-b', 'responses'))
        self.assertEqual(config.api_key, 'sk-parley-v1-test-only')
        self.assertFalse(self.app.session_state['api_permission_REAL'])
        self.assertNotIn('api_pending_config', self.app.session_state)
        self.assertNotIn('sk-parley-v1-test-only', str(Store(self.tmp.name, 'REAL').export()))
        next(b for b in self.app.button if b.label == 'Disconnect and forget session credentials').click().run()
        self.assertNotIn('api_config', self.app.session_state)
        self.assertNotIn('api_pending_config', self.app.session_state)
        self.assertNotIn('api_model_catalog', self.app.session_state)

    def test_pending_key_in_chat_is_blocked_without_an_applied_model(self):
        from sacfactory.storage import Store
        self.app.radio(key='page').set_value('Language model').run()
        next(b for b in self.app.button if b.label == 'Use MIT Parley preset').click().run()
        self.app.text_input(key='api_key_input').set_value('test-pending-private-key')
        next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
        self.app.radio(key='page').set_value('Experiment chat').run()
        self.app.chat_input[0].set_value('My key is test-pending-private-key').run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(Store(self.tmp.name, 'REAL').messages(), [])
        self.assertTrue(any('not saved or sent' in e.value for e in self.app.error))

    def test_endpoint_edit_clears_password_and_does_not_reuse_other_provider_key(self):
        self.configure_api()
        self.app.text_input(key='api_key_input').set_value('old-provider-widget-key').run()
        self.app.text_input(key='api_base_url').set_value('https://other.example.test/v1').run()
        self.assertEqual(self.app.text_input(key='api_key_input').value, '')
        self.assertFalse(self.app.session_state['api_permission_REAL'])
        with patch('sacfactory.providers.post_json') as network:
            next(b for b in self.app.button if b.label == 'Load available models').click().run()
            network.assert_not_called()
        self.assertTrue(any('API key is missing' in e.value for e in self.app.error))

    def test_research_disconnect_revokes_both_sharing_permissions(self):
        self.configure_api()
        self.app.radio(key='api_assistant_mode').set_value('Research conversation').run()
        next(w for w in self.app.checkbox if w.label.startswith('Share selected experimental records')).check().run()
        self.assertTrue(self.app.session_state['research_permission_REAL'])
        next(b for b in self.app.button if b.label == 'Disconnect and forget session credentials').click().run()
        self.assertFalse(self.app.session_state['api_permission_REAL'])
        self.assertFalse(self.app.session_state['research_permission_REAL'])
        self.assertNotIn('api_config', self.app.session_state)

    def test_responses_apply_selects_research_and_probe_does_not_share_records(self):
        from sacfactory.storage import Store
        self.configure_api()
        self.app.selectbox(key='api_format_choice').set_value('responses').run()
        next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
        self.assertEqual(self.app.session_state['assistant_mode_setting'], 'Research conversation')
        self.assertFalse(self.app.session_state['api_permission_REAL'])
        result = {'status': 'completed', 'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': 'OK'}]}]}
        with patch('sacfactory.providers.post_json', return_value=result) as network:
            next(b for b in self.app.button if b.label == 'Test connection (no experiment data)').click().run()
            request = network.call_args.args[0]
            import json
            self.assertEqual(json.loads(request.data)['input'], 'Reply with OK.')
            self.assertNotIn('REAL_ICP_001', request.data.decode())
            self.assertTrue(request.full_url.endswith('/responses'))
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any('Connection verified on /responses' in item.value for item in self.app.success))
        self.assertFalse(self.app.session_state['api_permission_REAL'])
        self.assertEqual(Store(self.tmp.name, 'REAL').messages(), [])

    def test_responses_task_mode_warning_and_probe_error_are_specific(self):
        from urllib.error import HTTPError
        self.configure_api()
        self.app.selectbox(key='api_format_choice').set_value('responses').run()
        next(b for b in self.app.button if b.label == 'Apply session settings').click().run()
        self.app.radio(key='api_assistant_mode').set_value('Task interpretation').run()
        self.assertTrue(any('Select Research conversation' in item.value for item in self.app.warning))
        with patch('sacfactory.providers.post_json', side_effect=HTTPError('PRIVATE_URL', 401, 'test-private-key', {}, None)):
            next(b for b in self.app.button if b.label == 'Test connection (no experiment data)').click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(any('HTTP 401, /responses' in item.value for item in self.app.error))
        self.assertFalse(any('test-private-key' in item.value for item in self.app.error))

    def test_unlimited_sharing_all_batches_navigation_and_namespace_isolation(self):
        self.app.radio(key='mode').set_value('Simulated demo · DEMO').run()
        self.app.radio(key='page').set_value('Language model').run()
        self.app.radio(key='api_assistant_mode').set_value('Research conversation').run()
        next(b for b in self.app.button if b.label == 'Select all records').click().run()
        self.assertEqual(len(self.app.session_state['research_samples_DEMO']), 1476)
        self.app.radio(key='page').set_value('Experiment chat').run()
        self.assertEqual(len(self.app.session_state['research_samples_DEMO']), 1476)
        next(b for b in self.app.button if b.label == 'Clear selection').click().run()
        self.assertEqual(self.app.session_state['research_samples_DEMO'], [])
        selector = next(w for w in self.app.multiselect if w.label == 'Batches to add')
        selector.set_value(selector.options[:12]).run()
        next(b for b in self.app.button if b.label == 'Add batches to selection').click().run()
        self.assertEqual(len(self.app.session_state['research_samples_DEMO']), 84)
        self.app.radio(key='page').set_value('Language model').run()
        self.assertEqual(len(self.app.session_state['research_samples_DEMO']), 84)
        self.app.radio(key='mode').set_value('Real experiments · REAL').run()
        self.assertEqual(self.app.session_state['research_samples_REAL'], [])
        self.assertTrue(all('DEMO_' not in i for i in next(w for w in self.app.multiselect if w.label == 'Shared experiment records').options))

    def test_research_prompt_buttons_use_api_and_local_shortcuts_remain_local(self):
        from sacfactory.dialogue import Answer
        from sacfactory.research_chat import ResearchAnswer
        self.configure_api()
        self.app.radio(key='api_assistant_mode').set_value('Research conversation').run()
        next(b for b in self.app.button if b.label == 'Select all records').click().run()
        next(w for w in self.app.checkbox if w.label.startswith('Share selected experimental records')).check().run()
        self.app.radio(key='page').set_value('Experiment chat').run()
        result = ResearchAnswer(Answer('Interpretation based on recorded evidence.', [], []), [], {'input_tokens': 1, 'output_tokens': 1}, 1)
        with patch('sacfactory.research_chat.ResearchAdapter.converse', return_value=result) as mock:
            next(b for b in self.app.button if b.label == 'Explain failures').click().run()
            mock.assert_called_once()
            self.assertIn('Analyze failures', mock.call_args.args[0])
            self.assertEqual(mock.call_args.args[2]['selected_record_count'], 1)
            next(b for b in self.app.button if b.label == 'Summary').click().run()
            self.assertEqual(mock.call_count, 1)
        self.assertEqual(len(self.app.exception), 0)

    def test_api_secret_in_chat_not_saved_or_sent(self):
        from sacfactory.storage import Store
        self.configure_api()
        self.app.radio(key="page").set_value("Experiment chat").run()
        with patch("sacfactory.language.CompatibleAdapter.interpret") as mock:
            self.app.chat_input[0].set_value("My key is test-private-key").run()
            mock.assert_not_called()
        self.assertEqual(Store(self.tmp.name, "REAL").messages(), [])
        self.assertTrue(any("not saved or sent" in e.value for e in self.app.error))

    def test_api_errors_no_proposals_and_disconnect(self):
        from sacfactory.language import LanguageError
        self.configure_api()
        self.app.radio(key="page").set_value("Experiment chat").run()
        with patch("sacfactory.language.CompatibleAdapter.interpret", side_effect=LanguageError("Authentication failed")):
            self.app.chat_input[0].set_value("Plan next batch").run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertIn("Authentication failed", "\n".join(x.value for x in self.app.markdown))
        self.app.radio(key="page").set_value("Language model").run()
        next(b for b in self.app.button if b.label == "Disconnect and forget session credentials").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertFalse(self.app.session_state["api_permission_REAL"])
        self.assertNotIn("api_config", self.app.session_state)

    def test_apply_new_endpoint_revokes_existing_permissions(self):
        self.configure_api()
        next(w for w in self.app.text_input if w.label == "API base URL").set_value("https://other.example.test/v1")
        next(b for b in self.app.button if b.label == "Apply session settings").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertFalse(self.app.session_state["api_permission_REAL"])
        self.assertFalse(self.app.session_state["api_permission_DEMO"])

    def test_periodic_selection_and_controlled_temperature_design(self):
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Next batch").run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        next(b for b in self.app.button if b.label == "Fe").click().run()
        self.assertEqual(self.app.session_state["plan_elements_DEMO"], ["Fe"])
        next(w for w in self.app.checkbox if w.label == "Include controlled cutoff temperature in the model").check()
        for label, value in (("Cutoff lower bound / °C", 900.), ("Cutoff upper bound / °C", 1100.)):
            next(w for w in self.app.number_input if w.label == label).set_value(value)
        next(w for w in self.app.checkbox if w.label == "This range is for simulation only").check()
        next(b for b in self.app.button if b.label == "Save operating range").click().run()
        self.assertEqual(len(self.app.exception), 0)
        next(b for b in self.app.button if b.label == "Generate batch preview").click().run()
        self.assertEqual(len(self.app.exception), 0)
        rows = self.app.session_state["preview_DEMO"]
        self.assertTrue(rows)
        self.assertTrue(all(900 <= r["recipe"]["pulse_limit_C"] <= 1100 for r in rows))
        self.assertTrue(all(r["recipe"]["temperature_C"] is None for r in rows))
        self.app.radio(key="page").set_value("Experiment chat").run()
        self.app.chat_input[0].set_value("Recommend 2 experiments").run()
        self.assertEqual(len(self.app.exception), 0)

    def test_fillable_batch_preview_save_and_reset(self):
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.radio if w.label == "Batch Sample ID numbering").set_value("Manual").run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        next(b for b in self.app.button if b.label == "Fe").click().run()
        next(w for w in self.app.number_input if w.label == "Samples in this reporting batch").set_value(1).run()
        for label, value in (("Synthesis batch ID", "TEST_ENTRY_BATCH"), ("Batch XRD source record", "DEMO spectrum fixture"), ("Batch XRD comparison criteria", "DEMO phase comparison"), ("XRD reviewer", "TESTER")):
            next(w for w in self.app.text_input if w.label == label).set_value(value)
        self.app.run()
        store = Store(self.tmp.name, "DEMO")
        methods = {m["method_id"]: m for m in store.all("method")}
        value = 0.25

        def filled(frame, *args, **kwargs):
            frame = frame.copy()
            if "concentration_mM" in frame.columns:
                frame.loc[0, "sample_id"] = "TEST_ENTRY_SAMPLE"
                frame.loc[0, "pulse_limit_C"] = "950"
                frame.loc[0, "temperature_C"] = "978"
                frame.loc[0, "voltage_V"] = 12.
                frame.loc[0, "current_A"] = 4.
                frame.loc[0, "icp_result"] = "PASS"
                frame.loc[0, "xrd_result"] = "PASS"
            elif str(kwargs.get('key', '')).startswith('element_grid_'):
                self.assertEqual(kwargs["column_config"]["value"]["label"], "W / wt%")
                self.assertEqual(set(frame["unit"]), {"wt%"})
                frame.loc[0, "sample_id"] = "TEST_ENTRY_SAMPLE"
                frame.loc[0, "value"] = str(value)
            return frame

        with patch("streamlit.data_editor", side_effect=filled):
            self.app.run()
            next(b for b in self.app.button if b.label == "Preview filled batch").click().run()
            self.assertEqual(len(self.app.error), 0)
            next(b for b in self.app.button if b.label == "Save filled batch").click().run()
        self.assertEqual(len(self.app.exception), 0)
        saved = store.get("sample", "TEST_ENTRY_SAMPLE")
        self.assertIsNotNone(saved)
        self.assertNotIn("qc", saved)
        self.assertEqual(screen(saved, methods)["status"], "PASS")
        self.assertEqual(saved["pulse_limit_C"], 950)
        self.assertEqual(saved["icp"][0]["value"], "0.25")
        self.assertEqual(saved["icp"][0]["unit"], "wt%")
        next(b for b in self.app.button if b.label == "Start an empty entry table").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertTrue(all(not r["value"] for r in self.app.session_state["entry_cells_DEMO"]["detail"]))

    def test_compact_workbook_upload_preview_commit(self):
        import openpyxl
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        root = Path(__file__).resolve().parents[1]
        book = openpyxl.load_workbook(root / "examples/compact_batch_template.xlsx")
        columns = {c.value: c.column for c in book["batch"][1]}
        for i in range(2, book["batch"].max_row + 1):
            book["batch"].cell(i, columns["pulse_limit_C"], 1000)
        cols = {c.value: c.column for c in book["settings"][1]}
        book["settings"].cell(2, cols["origin"], "DEMO")
        book["settings"].cell(2, cols["batch_id"], "DEMO_BATCH_FROM_XLSX")
        upload = io.BytesIO()
        book.save(upload)
        book.close()
        upload.name = "compact_demo.xlsx"
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Samples & data").run()

        def uploaded(label, *args, **kwargs):
            if label.startswith("Standard Excel"):
                return [upload] if kwargs["key"].endswith("_0") else []
            return None

        with patch("streamlit.file_uploader", side_effect=uploaded):
            self.app.run()
            next(b for b in self.app.button if b.label == "Commit import to local database").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.metric[0].value, "1482")
        store = Store(self.tmp.name, "DEMO")
        saved = store.get("sample", "SAC_001")
        self.assertEqual(saved["pulse_limit_C"], 1000)
        self.assertEqual(screen(saved, {})["status"], "PENDING")

    def test_per_sample_periodic_selection_preserves_entered_element_results(self):
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.number_input if w.label == "Samples in this reporting batch").set_value(2).run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        next(b for b in self.app.button if b.label == "Fe").click().run()

        def filled(frame, *args, **kwargs):
            frame = frame.copy()
            if "precursor_mM" in frame.columns: frame.loc[0, "value"] = "17.23"
            return frame

        with patch("streamlit.data_editor", side_effect=filled): self.app.run()
        next(w for w in self.app.selectbox if w.label == "Periodic-table selection applies to").set_value(2).run()
        next(b for b in self.app.button if b.label == "Mg").click().run()
        self.assertEqual(len(self.app.exception), 0)
        data = self.app.session_state["entry_cells_DEMO"]
        self.assertEqual(data["batch"][0]["target_elements"], "Fe")
        self.assertEqual(data["batch"][1]["target_elements"], "Mg;Fe")
        sid = data["batch"][0]["sample_id"]
        self.assertEqual(next(r["value"] for r in data["detail"] if r["sample_id"] == sid and r["element"] == "Fe"), "17.23")
        self.assertFalse(next(b for b in self.app.button if b.label == "Cd").disabled)

    def test_all_elements_and_direct_batch_pass_without_files_or_measured_values(self):
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.number_input if w.label == "Samples in this reporting batch").set_value(1).run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        for element in ("B", "N", "Cd", "Md"):
            control = next(b for b in self.app.button if b.label == element)
            self.assertFalse(control.disabled)
            control.click().run()
        next(w for w in self.app.selectbox if w.label == "ICP judgment for all rows").set_value("PASS")
        next(w for w in self.app.selectbox if w.label == "XRD judgment for all rows").set_value("PASS")
        next(b for b in self.app.button if b.label == "Apply judgments to all rows").click().run()
        self.assertEqual(self.app.session_state["entry_cells_DEMO"]["batch"][0]["icp_result"], "PASS")
        next(b for b in self.app.button if b.label == "Preview filled batch").click().run()
        self.assertEqual(len(self.app.error), 0)
        next(b for b in self.app.button if b.label == "Save filled batch").click().run()
        self.assertEqual(len(self.app.exception), 0)
        store = Store(self.tmp.name, "DEMO")
        s = store.get("sample", "SAC_0001")
        self.assertEqual(set(s["elements"]), {"B", "N", "Cd", "Md"})
        self.assertEqual(screen(s, {})["status"], "PASS")
        self.assertIsNone(s["temperature_C"])
        self.assertEqual(s["batch_id"], "UNKNOWN")
        self.assertTrue(all(not a["value"] for a in s["icp"]))
        self.app.radio(key="page").set_value("Next batch").run()
        self.assertEqual(len(self.app.exception), 0)

    def test_direct_concentration_ranges_save_and_recommend_without_multiplier_widgets(self):
        from sacfactory.storage import Store
        self.app.radio(key="mode").set_value("Simulated demo · DEMO").run()
        self.app.radio(key="page").set_value("Next batch").run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        for e in ("Fe", "Mg"):
            next(b for b in self.app.button if b.label == e).click().run()
        self.assertFalse(any("alpha" in w.label.lower() for w in self.app.number_input))

        def entered(frame, *args, **kwargs):
            frame = frame.copy()
            if "reference_mM" in frame.columns:
                for e, ref, lo, hi in (("Fe", .07, .04, .12), ("Mg", .11, .08, .16)):
                    frame.loc[frame.element == e, ["reference_mM", "min_mM", "max_mM"]] = [ref, lo, hi]
            return frame

        next(w for w in self.app.checkbox if w.label == "This range is for simulation only").check()
        with patch("streamlit.data_editor", side_effect=entered):
            next(b for b in self.app.button if b.label == "Save operating range").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        store = Store(self.tmp.name, "DEMO")
        ctx = store.get("meta", "plan_context")["context"]
        self.assertEqual(ctx["concentrations_mM"], {"Mg": .11, "Fe": .07})
        self.assertNotIn("alpha_min", store.get("meta", "space")["space"])
        next(b for b in self.app.button if b.label == "Generate batch preview").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        rows = self.app.session_state["preview_DEMO"]
        self.assertEqual(len(rows), 6)
        self.assertTrue(all(.04 <= r["recipe"]["concentrations_mM"]["Fe"] <= .12 for r in rows))
        self.assertTrue(any("concentration_mM" in frame.value.columns for frame in self.app.dataframe))

    def test_pending_recipe_registration_accepts_each_element_in_mM(self):
        from sacfactory.storage import Store
        self.app.radio(key="page").set_value("Samples & data").run()
        self.app.radio(key="register_preparation_REAL_unit").set_value("mM").run()
        next(w for w in self.app.radio if w.label == "Sample ID numbering").set_value("Manual").run()
        next(w for w in self.app.text_area if w.label == "Target elements, separated by semicolons").set_value("Fe;Mg").run()
        next(w for w in self.app.text_input if w.label == "New sample ID").set_value("DIRECT_001")

        def entered(frame, *args, **kwargs):
            frame = frame.copy()
            if list(frame.columns) == ["element", "precursor_mM"]:
                frame.loc[frame.element == "Fe", "precursor_mM"] = .073
                frame.loc[frame.element == "Mg", "precursor_mM"] = .119
            return frame

        with patch("streamlit.data_editor", side_effect=entered):
            next(b for b in self.app.button if b.label == "Register pending sample").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        self.assertEqual(Store(self.tmp.name, "REAL").get("sample", "DIRECT_001")["concentrations_mM"], {"Mg": .119, "Fe": .073})
        registered = Store(self.tmp.name, 'REAL').get('sample', 'DIRECT_001')['precursors']
        self.assertEqual(registered['Fe']['compound'], 'FeCl2')
        self.assertEqual(registered['Mg']['compound'], 'MgCl2\u00b76H2O')
        self.assertTrue(registered['Fe']['source'].startswith('PLANNED_CATALOG_ASSUMPTION:'))

    def test_empty_real_after_clear_all_pages_and_numbering(self):
        from sacfactory.storage import Store
        self.app.radio(key="page").set_value("Samples & data").run()
        next(b for b in self.app.button if b.label == "Clear all REAL data").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.metric[0].value, "0")
        self.assertEqual(Store(self.tmp.name, "REAL").next_sample_ids(), ["SAC_0001"])
        for page in ("Experiment chat", "Samples & data", "Next batch", "Materials & consumables", "Data analysis", "Language model", "About"):
            self.app.radio(key="page").set_value(page).run()
            self.assertEqual(len(self.app.exception), 0, page)
            self.assertEqual(self.app.metric[0].value, "0")

    def test_two_automatic_ui_batches_continue_at_eleven(self):
        from sacfactory.storage import Store
        self.app.radio(key="page").set_value("Samples & data").run()
        next(b for b in self.app.button if b.label == "Clear all REAL data").click().run()
        next(b for b in self.app.button if b.label == "Clear selection").click().run()
        next(b for b in self.app.button if b.label == "Fe").click().run()
        next(w for w in self.app.number_input if w.label == "Samples in this reporting batch").set_value(10).run()
        self.assertEqual(self.app.session_state["entry_cells_REAL"]["batch"][0]["sample_id"], "SAC_0001")
        self.assertEqual(self.app.session_state["entry_cells_REAL"]["batch"][-1]["sample_id"], "SAC_0010")
        next(w for w in self.app.text_input if w.label == "Actual independent synthesis batch ID").set_value("BATCH_1").run()
        next(b for b in self.app.button if b.label == "Preview filled batch").click().run()
        next(b for b in self.app.button if b.label == "Save filled batch").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        self.assertEqual(self.app.metric[0].value, "10")
        self.assertEqual(self.app.session_state["entry_cells_REAL"]["batch"][0]["sample_id"], "SAC_0011")
        self.assertTrue(all(not r["value"] for r in self.app.session_state["entry_cells_REAL"]["detail"]))
        next(w for w in self.app.text_input if w.label == "Actual independent synthesis batch ID").set_value("BATCH_2").run()
        next(b for b in self.app.button if b.label == "Preview filled batch").click().run()
        next(b for b in self.app.button if b.label == "Save filled batch").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(len(self.app.error), 0)
        db = Store(self.tmp.name, "REAL")
        self.assertEqual(len(db.all("sample")), 20)
        self.assertEqual(db.get("sample", "SAC_0010")["batch_id"], "BATCH_1")
        self.assertEqual(db.get("sample", "SAC_0011")["batch_id"], "BATCH_2")
        self.assertEqual(db.next_sample_ids(), ["SAC_0021"])

    def test_delete_selected_sample_ui_does_not_reseed_it(self):
        from sacfactory.storage import Store
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.multiselect if w.label == "Samples to delete").set_value(["REAL_ICP_001"]).run()
        next(b for b in self.app.button if b.label == "Delete selected samples").click().run()
        self.assertEqual(len(self.app.exception), 0)
        self.assertEqual(self.app.metric[0].value, "0")
        self.app.run()
        self.assertIsNone(Store(self.tmp.name, "REAL").get("sample", "REAL_ICP_001"))

    def test_existing_sample_direct_judgment_update_without_evidence_files(self):
        from sacfactory.storage import Store
        from sacfactory.domain import screen
        self.app.radio(key="page").set_value("Samples & data").run()
        next(w for w in self.app.selectbox if w.label == "ICP judgment").set_value("PASS")
        next(w for w in self.app.selectbox if w.label == "XRD relevant-phase criterion").set_value("PASS")
        next(b for b in self.app.button if b.label == "Save evidence revision").click().run()
        self.assertEqual(len(self.app.error), 0)
        self.assertEqual(len(self.app.exception), 0)
        s = Store(self.tmp.name, "REAL").get("sample", "REAL_ICP_001")
        self.assertEqual(screen(s, {})["status"], "PASS")
        self.assertEqual(s["icp_review"]["state"], "PASS")
        self.assertEqual(len(s["icp"]), 51)


if __name__ == "__main__": unittest.main()
