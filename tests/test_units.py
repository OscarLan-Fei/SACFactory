from sacfactory.bootstrap import load_demo_seed
"""Mass-percent recording, legacy conversion and migration regression tests."""
import copy
import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

import openpyxl

from sacfactory.units import normalize_assay, normalize_method, to_wt_percent
from sacfactory.storage import Store
from sacfactory.domain import screen, validate_sample
from sacfactory.batching import blank_batch, batch_rows_to_standard
from sacfactory.importers import samples_from_rows, sample_tables, csv_bytes, read_table, ICP_COLUMNS, methods_from_rows, METHOD_COLUMNS
from test_science import sample, method


class MassPercentTests(unittest.TestCase):
    def test_percent_points_not_fraction(self):
        for unit in ("wt%", "wt.%", "W%", "%"):
            self.assertEqual(normalize_assay({"value": "0.25", "unit": unit})["value"], "0.25")
        self.assertEqual(normalize_assay({"W": "0.25"}), {"value": "0.25", "unit": "wt%"})

    def test_legacy_values_and_censoring_convert_explicitly(self):
        for value, expected in (("2500", "0.25"), ("<50", "<0.005"), ("≤100", "≤0.01"), ("0", "0"), (2500.0, .25)):
            self.assertEqual(to_wt_percent(value, "ug/g"), expected)
        for token in ("", "ND", "BLQ", "<LOQ", "N.D."):
            self.assertEqual(to_wt_percent(token, "ug/g"), token)

    def test_unknown_units_and_conflicting_headers_are_rejected(self):
        for row in ({"value": "2", "unit": "ug/L"}, {"W": "0.25", "unit": "ug/g"}, {"W": "0.25", "value": "2500"}, {"W": "0.25", "w_wt_percent": "0.25"}):
            with self.assertRaises(ValueError):
                normalize_assay(row)

    def test_standard_w_and_legacy_imports_agree(self):
        s = sample()
        s["icp"][0]["value"] = "0.25"
        ss, ii = sample_tables([s])
        exported = read_table(csv_bytes(ii, ICP_COLUMNS), "icp.csv")
        self.assertEqual(exported[0]["w_wt_percent"], "0.25")
        self.assertNotIn("value", exported[0])
        restored = samples_from_rows(ss, exported)[0]
        self.assertEqual(restored["icp"][0]["value"], "0.25")
        legacy = copy.deepcopy(ii)
        legacy[0].update(value="2500", unit="ug/g")
        legacy = read_table(csv_bytes(legacy, ["value" if c == "w_wt_percent" else c for c in ICP_COLUMNS]), "legacy_icp.csv")
        self.assertEqual(samples_from_rows(ss, legacy)[0]["icp"], restored["icp"])

    def test_inline_and_detail_content_use_wt_percent(self):
        rows, details = blank_batch(["SAC_0001"], ["Fe"])
        settings = {"target_elements": "Fe", "origin": "REAL"}
        rows[0].update(icp_result="Fe:0.25", xrd_result="PASS")
        ss, ii = batch_rows_to_standard(rows, details, settings)
        self.assertEqual(ii[0]["unit"], "wt%")
        self.assertEqual(samples_from_rows(ss, ii)[0]["icp"][0]["value"], "0.25")
        rows[0]["icp_result"] = "PASS"
        details[0].update(value="2500", unit="ug/g")
        ss, ii = batch_rows_to_standard(rows, details, settings)
        self.assertEqual(samples_from_rows(ss, ii)[0]["icp"][0]["value"], "0.25")

    def test_method_limits_use_percent_points_and_read_legacy_headers(self):
        row = {"method_id": "M1", "origin": "REAL", "instrument": "ICP-AES", "rule": "LOQ", "element": "Fe", "lod_wt_percent": "0.001", "loq_wt_percent": "0.005", "confirmed": "1", "source": "Test", "version": "v1"}
        current = methods_from_rows([row])[0]
        self.assertEqual(current["unit"], "wt%")
        self.assertEqual(current["limits"]["Fe"]["loq"], .005)
        legacy = {k: v for k, v in row.items() if k not in {"lod_wt_percent", "loq_wt_percent"}}
        legacy.update(lod_ug_g="10", loq_ug_g="50")
        self.assertEqual(methods_from_rows([legacy])[0], current)

    def test_numeric_boundaries_and_labels_are_invariant_after_conversion(self):
        legacy_method = method()["M1"]
        legacy_method.pop("unit")
        converted_method = normalize_method(legacy_method, legacy=True)
        for value in ("0", "1", "10", "12", "<10", "≤10", "ND", "BLQ"):
            old = sample()
            old["icp"][0].update(value=value, unit="ug/g", flag="ND" if value == "ND" else "BLQ" if value == "BLQ" else "MEASURED")
            old_result = screen(old, {"M1": legacy_method})
            converted = validate_sample(copy.deepcopy(old))
            new_result = screen(converted, {"M1": converted_method})
            self.assertEqual((old_result["status"], old_result["label"]), (new_result["status"], new_result["label"]))
            if old_result["weakest"]:
                self.assertAlmostEqual(old_result["weakest"][1], new_result["weakest"][1])

    def test_atomic_idempotent_database_migration_retains_original_audit(self):
        with tempfile.TemporaryDirectory() as folder:
            db = Store(folder, "DEMO")
            old = sample(origin="DEMO")
            old["icp"][0].update(value="2500", unit="ug/g")
            m = method("DEMO")["M1"]
            m.pop("unit")
            m["limits"]["Fe"] = {"lod": 10, "loq": 50}
            with db.connect() as c:
                c.execute("DELETE FROM items WHERE id='icp-mass-percent-v1'")
                for kind, ident, obj in (("sample", old["sample_id"], old), ("method", "M1", m), ("meta", "plan_context", {"id": "plan_context", "origin": "DEMO", "context": old})):
                    c.execute("INSERT INTO items VALUES(?,?,?)", (kind, ident, json.dumps(obj)))
            migrated = Store(folder, "DEMO")
            self.assertEqual(migrated.get("sample", old["sample_id"])["icp"][0]["value"], "0.25")
            self.assertEqual(migrated.get("method", "M1")["limits"]["Fe"]["loq"], .005)
            self.assertEqual(migrated.get("meta", "plan_context")["context"]["icp"][0]["unit"], "wt%")
            before = migrated.export()
            after = Store(folder, "DEMO").export()
            self.assertEqual(before, after)
            revisions = [r for r in after["audit"] if r["kind"] == "sample"]
            self.assertIn('"unit": "ug/g"', revisions[0]["before_json"])
            self.assertIn('"unit": "wt%"', revisions[0]["after_json"])
            self.assertFalse((Path(folder) / "backups").exists())

    def test_bad_legacy_unit_rolls_back_migration(self):
        with tempfile.TemporaryDirectory() as folder:
            db = Store(folder, "DEMO")
            with db.connect() as c:
                c.execute("DELETE FROM items WHERE id='icp-mass-percent-v1'")
                for i, unit in enumerate(("ug/g", "ug/L")):
                    s = sample(i, origin="DEMO")
                    s["icp"][0].update(value="2500", unit=unit)
                    c.execute("INSERT INTO items VALUES('sample',?,?)", (s["sample_id"], json.dumps(s)))
            with self.assertRaises(ValueError):
                Store(folder, "DEMO")
            self.assertEqual(db.get("sample", "S0")["icp"][0]["unit"], "ug/g")
            self.assertIsNone(db.get("meta", "icp-mass-percent-v1"))

    def test_templates_declare_mass_percent_and_do_not_format_as_fraction(self):
        root = Path(__file__).resolve().parents[1]
        for name in ("compact_batch_template.xlsx", "bulk_import_template.xlsx"):
            book = openpyxl.load_workbook(root / "examples" / name)
            try:
                cols = {c.value: c.column for c in book["icp"][1]}
                self.assertIn("w_wt_percent", cols)
                self.assertNotIn("value", cols)
                for i in range(2, book["icp"].max_row + 1):
                    self.assertEqual(book["icp"].cell(i, cols["unit"]).value, "wt%")
                    self.assertNotIn("%", book["icp"].cell(i, cols["w_wt_percent"]).number_format)
            finally:
                book.close()

    def test_seed_is_canonical_and_keeps_expected_physical_scale(self):
        root = Path(__file__).resolve().parents[1]
        seed = load_demo_seed(root)
        self.assertEqual(len(seed["samples"]), 1476)
        self.assertTrue(all(a["unit"] == "wt%" for s in seed["samples"] for a in s["icp"]))
        values = [float(a['value']) for s in seed['samples'] for a in s['icp'] if a['flag'] == 'MEASURED']
        self.assertTrue(all(0 <= value <= 100 for value in values))
        self.assertEqual(seed['demo_version'], 'synthetic-physics-linked-campaign-v6')
        self.assertTrue(all(m["unit"] == "wt%" for m in seed["methods"]))


if __name__ == "__main__":
    unittest.main()
