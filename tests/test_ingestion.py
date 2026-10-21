"""Exercise the report-feedback cycle with labeled DEMO fixtures only."""
import copy
import io
import tempfile
import unittest

import openpyxl

from sacfactory.domain import screen, recipe_key, RULE_VERSION
from sacfactory.importers import recommendation_workbook, read_table, csv_bytes, METHOD_COLUMNS
from sacfactory.ingestion import preview_import, commit_import
from sacfactory.planning import fit_model, recommend
from sacfactory.storage import Store
from test_science import sample, method


def measured_workbook(proposals, values="15"):
    """Fabricated test results, kept strictly in DEMO and temporary databases."""
    book = openpyxl.load_workbook(io.BytesIO(recommendation_workbook(proposals)))
    sample_sheet = book["samples"]
    columns = {c.value: c.column for c in sample_sheet[1]}
    for i in range(2, sample_sheet.max_row + 1):
        updates = {"batch_id": f"DEMO_RETURN_BATCH_{i}", "xrd_state": "PASS", "xrd_source": "DEMO spectrum fixture", "xrd_criteria": "DEMO reference comparison", "xrd_reviewer": "TEST_FIXTURE"}
        for key, value in updates.items():
            sample_sheet.cell(i, columns[key], value)
    icp_sheet = book["icp"]
    columns = {c.value: c.column for c in icp_sheet[1]}
    for i in range(2, icp_sheet.max_row + 1):
        updates = {"w_wt_percent": values, "flag": "MEASURED", "method_id": "M1", "source": "DEMO ICP fixture"}
        for key, value in updates.items():
            icp_sheet.cell(i, columns[key], value)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def method_csv(m):
    return csv_bytes([{"method_id": m["method_id"], "origin": m["origin"], "instrument": m["instrument"], "rule": m["rule"], "element": e,
        "lod_wt_percent": limits["lod"], "loq_wt_percent": limits["loq"], "confirmed": "1" if m["confirmed"] else "0", "source": m["source"], "version": m["version"]}
        for e, limits in m["limits"].items()], METHOD_COLUMNS)


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Store(self.tmp.name, "DEMO")
        self.methods = method("DEMO")
        self.db.put_many("method", list(self.methods.values()), "method_id")
        self.s = sample(origin="DEMO")
        self.db.put_many("sample", [self.s], "sample_id")
        self.space = {"alpha_min": .5, "alpha_max": 2, "pulses": [1, 2], "max_total_mM": 1, "verified": True, "baseline_repeat": False}

    def tearDown(self):
        self.tmp.cleanup()

    def proposals(self, count=3):
        state = fit_model(self.db.all("sample"), self.methods, self.s, "DEMO")
        return recommend(state, self.db.all("sample"), self.s, self.space, "DEMO", count)[0]

    def test_full_reserve_export_measure_import_recommend_cycle(self):
        proposed = self.proposals()
        self.db.put_many("recommendation", proposed, "recommendation_id")
        payload = measured_workbook(proposed)
        preview = preview_import([("demo_results.xlsx", payload)], self.db)
        self.assertTrue(all(r["status"] == "Candidate screening passed" for r in preview.rows))
        self.assertEqual(len(preview.queue_updates), 3)
        result = commit_import(preview, self.db)
        self.assertEqual(result, {"samples": 3, "methods": 0, "completed_tasks": 3})
        self.assertTrue(all(r["state"] == "COMPLETED" for r in self.db.all("recommendation")))
        for row in self.db.all("sample"):
            if row["sample_id"] != self.s["sample_id"]:
                self.assertEqual(screen(row, self.methods)["label"], 1)
                self.assertTrue(row["import_sources"][0]["sha256"])
        new = self.proposals()
        self.assertFalse({recipe_key(r["recipe"]) for r in proposed} & {recipe_key(r["recipe"]) for r in new})

    def test_report_revision_changes_label_preserves_original(self):
        proposed = self.proposals(1)
        first = measured_workbook(proposed, "15")
        commit_import(preview_import([("pass.xlsx", first)], self.db), self.db)
        failed = measured_workbook(proposed, "1")
        with self.assertRaises(ValueError): preview_import([("fail.xlsx", failed)], self.db)
        result = commit_import(preview_import([("fail.xlsx", failed)], self.db, overwrite=True), self.db)
        s = self.db.get("sample", proposed[0]["recommendation_id"])
        self.assertEqual(screen(s, self.methods)["status"], "ICP_FAIL")
        self.assertEqual(len(s["import_sources"]), 2)
        revisions = [a for a in self.db.export()["audit"] if a["id"] == s["sample_id"]]
        self.assertEqual(len(revisions), 2)
        self.assertIn('"value": "15"', revisions[1]["before_json"])

    def test_method_revision_is_explicit_and_reports_affected_samples(self):
        m = copy.deepcopy(self.methods["M1"])
        m["version"] = "v2"
        m["limits"]["Fe"]["loq"] = 20
        files = [("methods.csv", method_csv(m))]
        with self.assertRaises(ValueError): preview_import(files, self.db)
        preview = preview_import(files, self.db, overwrite=True)
        self.assertEqual(preview.method_changes[0]["affected_existing_samples"], 1)
        old_label = screen(self.s, self.methods)
        commit_import(preview, self.db)
        revised = {m["method_id"]: m for m in self.db.all("method")}
        new_label = screen(self.s, revised)
        self.assertEqual(new_label["label"], 0)
        self.assertNotEqual(old_label["method_fingerprints"], new_label["method_fingerprints"])

    def test_identical_method_does_not_require_revision_or_create_update(self):
        preview = preview_import([("methods.csv", method_csv(self.methods["M1"]))], self.db)
        self.assertEqual(preview.method_changes, [])
        self.assertEqual(commit_import(preview, self.db)["methods"], 0)

    def test_pending_workbook_does_not_complete_task(self):
        proposed = self.proposals(1)
        self.db.put_many("recommendation", proposed, "recommendation_id")
        preview = preview_import([("pending.xlsx", recommendation_workbook(proposed))], self.db)
        self.assertEqual(preview.queue_updates, [])
        commit_import(preview, self.db)
        self.assertEqual(self.db.all("recommendation")[0]["state"], "OPEN")

    def test_altered_reserved_recipe_is_rejected(self):
        proposed = self.proposals(1)
        self.db.put_many("recommendation", proposed, "recommendation_id")
        altered = copy.deepcopy(proposed)
        altered[0]["recipe"]["pulse_count"] += 1
        with self.assertRaises(ValueError): preview_import([("altered.xlsx", measured_workbook(altered))], self.db)

    def test_real_namespace_rejects_demo_files(self):
        proposed = self.proposals(1)
        db = Store(self.tmp.name, "REAL")
        with self.assertRaises(ValueError): preview_import([("demo.xlsx", measured_workbook(proposed))], db)

    def test_duplicate_headers_rejected(self):
        with self.assertRaises(ValueError): read_table(b"value,value\n12,1\n", "invalid.csv")

    def test_blq_with_no_numeric_value_fails_loq_but_not_forced_lod(self):
        s = copy.deepcopy(self.s)
        s["icp"][0].update(flag="BLQ", value="")
        self.assertEqual(screen(s, self.methods)["status"], "ICP_FAIL")
        s["icp"][0].update(flag="MEASURED", value="<LOQ")
        self.assertEqual(screen(s, self.methods)["status"], "ICP_FAIL")
        limits = copy.deepcopy(self.methods)
        limits["M1"]["rule"] = "LOD"
        self.assertEqual(screen(s, limits)["status"], "PENDING")
        s["icp"][0].update(flag="BLQ", value="5")
        self.assertEqual(screen(s, limits)["status"], "PASS")


if __name__ == "__main__": unittest.main()
